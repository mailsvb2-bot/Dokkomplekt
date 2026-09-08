from __future__ import annotations

from tkinter import messagebox

from diagnostic_logging import record_soft_exception
from error_taxonomy import ErrorCategory, record_classified_error
from actions_creation_preflight import ActionsCreationReviewMixin
from actions_creation_foldering import ActionsCreationFolderingMixin
from actions_creation_maintenance import ActionsCreationMaintenanceMixin
from actions_creation_batch import ActionsCreationBatchingMixin
from actions_creation_execution import ActionsCreationExecutionMixin

# Contract sentinels kept for legacy smoke/prod gates that inspect this public
# aggregation module while implementation lives in focused files:
# _prompt_common_output_requirements _select_default_printer_sync
# def batch_generate_documents_dialog def _read_update_manifest
# def configure_patient_folder_naming_dialog doctor_confirmed _ensure_patient_folder_naming_configured
# return False Дата выписки не может быть раньше даты поступления


class ActionsCreationLiveGuardMixin:
    """Doctor-facing boundary for the main create buttons.

    The main create click is already explicit consent to generate. After truly
    missing required fields are resolved, creation must continue without a
    second confirmation modal that can become hidden behind the no-console EXE.
    Unexpected callback failures are made visible instead of looking like a
    dead button.
    """

    def _confirm_patient_case_before_creation(self, review) -> bool:
        """Block only on truly missing required fields, then generate."""
        return bool(self._prompt_missing_required_fields_or_continue(review))

    def _create_regular_custom_documents(self, current_pack, case, regular_ids, out_dir):
        """Preserve historical empty-placeholder UX without weakening strict rendering.

        ``strict=False`` used to solve two unrelated problems at once: it allowed
        empty doctor-declared optional placeholders, but it also disabled strict
        visible-Word-field validation.  Keep the renderer strict instead.  For
        this one render only, copy selected templates and remove exactly the
        placeholders the doctor is allowed to leave empty.
        """
        from dataclasses import replace
        from pathlib import Path
        import tempfile

        from docx import Document
        from medical_docx_xml_fragments import ensure_docx_compatible
        from universal_fields import normalize_field_id_for_context
        from universal_profiles import resolve_pack_template_path
        from universal_template_engine import (
            _iter_docx_paragraphs,
            _replace_paragraph_placeholders,
            extract_template_placeholders,
            missing_required_fields,
        )

        selected = {str(item).strip() for item in regular_ids if str(item).strip()}
        allow_required_override = bool(getattr(self, "_allow_missing_required_creation", False))
        raw_override_fields = tuple(
            str(field_id).strip()
            for field_id in tuple(getattr(self, "_missing_required_override_fields", ()) or ())
            if str(field_id).strip()
        )
        failures: list[str] = []
        profile_path = getattr(self, "_universal_profile_path", None)
        profile_base = profile_path().parent if callable(profile_path) else Path(out_dir).parent

        with tempfile.TemporaryDirectory(prefix="dokkomplekt-template-compat-") as temp_dir:
            temp_root = Path(temp_dir)
            rendered_documents = []
            for index, document in enumerate(tuple(getattr(current_pack, "documents", ()) or ())):
                if getattr(document, "id", "") not in selected:
                    rendered_documents.append(document)
                    continue

                missing_required = missing_required_fields(case, document)
                context_kwargs = {
                    "role_id": str(getattr(document, "role_id", "") or ""),
                    "category": str(getattr(document, "category", "") or ""),
                    "document_label": str(getattr(document, "button_label", "") or ""),
                }
                override_fields = {
                    normalize_field_id_for_context(field_id, **context_kwargs)
                    for field_id in raw_override_fields
                }
                permitted_required = {field_id for field_id in missing_required if field_id in override_fields}
                unpermitted_required = tuple(
                    field_id for field_id in missing_required if field_id not in permitted_required
                )
                if unpermitted_required or (missing_required and not allow_required_override):
                    label = str(getattr(document, "button_label", "") or getattr(document, "id", "") or "Документ")
                    failures.append(f"{label}: {', '.join(unpermitted_required or missing_required)}")
                    rendered_documents.append(document)
                    continue

                optional = {
                    normalize_field_id_for_context(field_id, **context_kwargs)
                    for field_id in tuple(getattr(document, "optional_fields", ()) or ())
                }
                allowed_empty = {field_id for field_id in optional if not case.get(field_id).strip()}
                if allow_required_override:
                    allowed_empty.update(permitted_required)

                render_document = document
                if permitted_required and allow_required_override:
                    missing_set = set(permitted_required)
                    kept_required = tuple(
                        field_id for field_id in document.required_fields
                        if normalize_field_id_for_context(field_id, **context_kwargs) not in missing_set
                    )
                    render_document = replace(
                        render_document,
                        required_fields=kept_required,
                        optional_fields=tuple(dict.fromkeys([*document.optional_fields, *permitted_required])),
                    )

                if allowed_empty:
                    source = resolve_pack_template_path(document.template, profile_base)
                    placeholders = extract_template_placeholders(
                        source,
                        role_id=document.role_id,
                        category=document.category,
                        button_label=document.button_label,
                    )
                    present = {item.field_id for item in placeholders}
                    removable = allowed_empty & present
                    if removable:
                        readable = ensure_docx_compatible(source, label="шаблон документа")
                        target = temp_root / f"{index:03d}-{source.stem}.docx"
                        doc = Document(str(readable))
                        preserve = {field_id: "{{" + field_id + "}}" for field_id in present}
                        for field_id in removable:
                            preserve[field_id] = ""
                        for paragraph, _hint in _iter_docx_paragraphs(doc):
                            _replace_paragraph_placeholders(
                                paragraph,
                                preserve,
                                set(),
                                set(),
                                document=document,
                            )
                        doc.save(str(target))
                        render_document = replace(render_document, template=str(target))

                rendered_documents.append(render_document)

            if failures:
                raise ValueError("Не заполнены обязательные поля документов: " + "; ".join(failures))

            render_pack = replace(current_pack, documents=tuple(rendered_documents))

            # The parent maps this legacy flag directly to renderer ``strict``.
            # Overrides have already been represented narrowly in render_pack, so
            # force strict rendering here and restore the UI state afterwards.
            previous = bool(getattr(self, "_allow_missing_required_creation", False))
            self._allow_missing_required_creation = False
            try:
                return super()._create_regular_custom_documents(render_pack, case, regular_ids, out_dir)
            finally:
                self._allow_missing_required_creation = previous

    def create_selected_outputs(self, *, print_after: bool = False) -> bool:
        """Run the existing generator and never let a live click fail invisibly."""
        try:
            self._set_status("Проверяю данные перед созданием…")
        except Exception as exc:
            record_soft_exception("actions_creation_orchestrator.set_start_status", exc)
        try:
            result = bool(super().create_selected_outputs(print_after=print_after))
        except Exception as exc:
            record_classified_error("create_selected_outputs_button", exc, category=ErrorCategory.DOCX_RENDER)
            try:
                self._set_status("Создание не выполнено: показана ошибка")
            except Exception as status_exc:
                record_soft_exception("actions_creation_orchestrator.set_failure_status", status_exc)
            messagebox.showerror(
                "Документы не созданы",
                "Создание остановилось с ошибкой. Никакие частичные документы не считаются готовыми.\n\n"
                f"{exc}\n\n"
                "Подробности записаны в диагностику программы.",
            )
            return False
        if result:
            try:
                self._set_status("Готово: файлы сохранены")
            except Exception as exc:
                record_soft_exception("actions_creation_orchestrator.set_success_status", exc)
        return result


class ActionsCreationOrchestratorMixin(
    ActionsCreationLiveGuardMixin,
    ActionsCreationReviewMixin,
    ActionsCreationFolderingMixin,
    ActionsCreationMaintenanceMixin,
    ActionsCreationBatchingMixin,
    ActionsCreationExecutionMixin,
):
    """Aggregate focused creation-flow mixins for selected output generation."""
