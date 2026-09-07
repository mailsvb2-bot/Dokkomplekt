"""Template-preserving renderer for doctor-owned custom diary buttons.

Standard Dates+Texts diaries intentionally stay on the text-document engine.
This module is only for a doctor-owned block-03 diary template: the selected
Word file remains the owner of layout, headers, footers, tables and signatures.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Sequence

from docx import Document

from diary_dates import parse_full_datetime, parse_optional_discharge_date
from diary_gender import adapt_text_to_patient_gender, detect_gender_from_patient_name
from diary_schedule import DiaryScheduleSpec, planned_diary_datetimes
from diary_text_parser import clean_status_text, is_signature_paragraph_text, remove_examinee_words
from diary_writer_apply import NEUTRAL_FINAL_DIARY_TEXT
from medical_docx_xml_fragments import ensure_docx_compatible
from medical_formatting import available_path
from universal_template_engine import (
    _iter_docx_paragraphs,
    _replace_paragraph_placeholders,
    build_render_context,
    missing_required_fields,
    remove_forbidden_hospitalization_phrase_from_document,
    render_output_name,
)


@dataclass(frozen=True)
class CustomDiaryTemplateRenderResult:
    path: Path
    filled_rows: int
    final_rows_filled: int


def _norm(value: str) -> str:
    return " ".join(str(value or "").replace("ё", "е").casefold().split())


def _header_columns(table) -> tuple[int, int | None, int | None, int | None, int] | None:
    """Return diary/day/month/hospitalization columns and header row index."""

    for row_index, row in enumerate(table.rows[:6]):
        texts = [_norm(cell.text) for cell in row.cells]
        diary_col = next((i for i, text in enumerate(texts) if "дневник" in text or "наблюден" in text), None)
        if diary_col is None:
            continue
        day_col = next((i for i, text in enumerate(texts) if "число" in text or text == "дата" or "дата наблю" in text), None)
        month_col = next((i for i, text in enumerate(texts) if "месяц" in text), None)
        hosp_col = next((i for i, text in enumerate(texts) if "госпит" in text and "день" in text), None)
        return diary_col, day_col, month_col, hosp_col, row_index
    return None


def _best_diary_table(doc: Document):
    for table in doc.tables:
        columns = _header_columns(table)
        if columns is not None:
            return table, columns
    return None, None


def _replace_content_paragraph_preserving_style(cell, text: str) -> None:
    """Replace generated content while leaving signature paragraphs/styles alive."""

    content_paragraphs = [p for p in cell.paragraphs if not is_signature_paragraph_text(p.text)]
    if not content_paragraphs:
        paragraph = cell.add_paragraph("")
        content_paragraphs = [paragraph]
    first = content_paragraphs[0]
    if first.runs:
        first.runs[0].text = str(text or "")
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.text = str(text or "")
    for paragraph in content_paragraphs[1:]:
        if paragraph.runs:
            for run in paragraph.runs:
                run.text = ""
        else:
            paragraph.text = ""


def _set_simple_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph("")
    if paragraph.runs:
        paragraph.runs[0].text = str(text or "")
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = str(text or "")
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def _status_entries(
    statuses: Sequence[str],
    moments: Sequence[datetime],
    *,
    patient_name: str,
    repeat_statuses: bool,
) -> list[tuple[datetime, str, bool]]:
    gender = detect_gender_from_patient_name(patient_name)
    result: list[tuple[datetime, str, bool]] = []
    index = 0
    for moment in moments:
        if not statuses:
            break
        if index >= len(statuses):
            if not repeat_statuses:
                break
            index = 0
        raw = statuses[index]
        index += 1
        adapted, _changed = adapt_text_to_patient_gender(raw, gender)
        result.append((moment, clean_status_text(adapted), False))
    return result


def _with_final_entry(
    entries: list[tuple[datetime, str, bool]],
    *,
    admission: datetime,
    discharge,
    force_final_diary: bool,
    patient_name: str,
) -> list[tuple[datetime, str, bool]]:
    if not force_final_diary or discharge is None:
        return entries
    gender = detect_gender_from_patient_name(patient_name)
    final_text, _changed = adapt_text_to_patient_gender(NEUTRAL_FINAL_DIARY_TEXT, gender)
    final_text = remove_examinee_words(final_text)
    final_moment = datetime.combine(discharge, admission.time())
    # A neutral discharge conclusion owns only its own final record; regular
    # observations earlier on the discharge day remain valid in hourly mode.
    if any(is_final and moment.date() == discharge for moment, _text, is_final in entries):
        return entries
    return [*entries, (final_moment, final_text, True)]


def _format_entry(moment: datetime, text: str, *, hourly: bool) -> str:
    prefix = moment.strftime("%d.%m.%y %H:%M") if hourly else moment.strftime("%d.%m.%y")
    return f"{prefix} {text}".rstrip()


def render_custom_diary_template(
    *,
    template_path: str | Path,
    output_dir: str | Path,
    case,
    document,
    schedule: DiaryScheduleSpec,
    statuses: Sequence[str],
    patient_name: str,
    admission_value: str,
    discharge_value: str,
    repeat_statuses: bool,
    force_final_diary: bool,
    output_language: str = "auto",
    spellcheck_enabled: bool = True,
) -> CustomDiaryTemplateRenderResult:
    """Render one custom diary without replacing the doctor's Word layout."""

    missing = missing_required_fields(case, document)
    if missing:
        raise ValueError("Не заполнены обязательные поля для дневника: " + ", ".join(missing))
    if not statuses:
        raise ValueError("В выбранном источнике не найдено текстов дневников.")

    admission = parse_full_datetime(admission_value)
    discharge = parse_optional_discharge_date(discharge_value)
    if discharge is not None and discharge < admission.date():
        raise ValueError("Дата выписки не может быть раньше даты поступления.")

    source = ensure_docx_compatible(template_path, label="шаблон дневников")
    doc = Document(str(source))
    table, columns = _best_diary_table(doc)
    if table is None or columns is None:
        raise ValueError(
            "В шаблоне дневников не найдена таблица с колонкой «Дневник/Наблюдение». "
            "Программа не будет заменять структуру вашего Word-шаблона новым документом."
        )

    diary_col, day_col, month_col, hosp_col, header_index = columns
    data_rows = list(table.rows[header_index + 1 :])
    if not data_rows:
        raise ValueError("В шаблоне дневников нет строк для заполнения после заголовка таблицы.")

    # Request one record beyond capacity so overflow becomes a visible error,
    # never silent truncation of the hospitalization episode.
    moments = list(planned_diary_datetimes(admission, schedule, limit=len(data_rows) + 1))
    if discharge is not None:
        moments = [moment for moment in moments if moment.date() <= discharge]
    entries = _status_entries(statuses, moments, patient_name=patient_name, repeat_statuses=repeat_statuses)
    entries = _with_final_entry(
        entries,
        admission=admission,
        discharge=discharge,
        force_final_diary=force_final_diary,
        patient_name=patient_name,
    )
    entries.sort(key=lambda item: (item[0], 1 if item[2] else 0))
    if len(entries) > len(data_rows):
        raise ValueError(
            f"В шаблоне дневников недостаточно строк: нужно {len(entries)}, доступно {len(data_rows)}. "
            "Увеличьте число строк в вашем Word-шаблоне или измените ритм дневников."
        )

    # Build the normal semantic context, then replace diary placeholders with the
    # actual observation text rather than a list of dates.
    context = build_render_context(case, document, output_language=output_language, spellcheck_enabled=spellcheck_enabled)
    hourly = schedule.mode == "hourly"
    rendered_entries = [_format_entry(moment, text, hourly=hourly) for moment, text, _final in entries]
    context["diary.entries"] = "\n".join(rendered_entries)
    context["diary.dates"] = "\n".join(
        moment.strftime("%d.%m.%Y %H:%M" if hourly else "%d.%m.%Y") for moment, _text, _final in entries
    )

    replaced: set[str] = set()
    missing_seen: set[str] = set()
    for paragraph, _hint in _iter_docx_paragraphs(doc):
        _replace_paragraph_placeholders(paragraph, context, replaced, missing_seen, document=document)
    blocking_missing = {field_id for field_id in missing_seen if field_id in set(document.required_fields)}
    if blocking_missing:
        raise ValueError("Не заполнены обязательные поля шаблона дневников: " + ", ".join(sorted(blocking_missing)))

    for row_index, row in enumerate(data_rows):
        if row_index < len(entries):
            moment, text, _is_final = entries[row_index]
            diary_text = f"{moment:%H:%M} {text}".rstrip() if hourly else text
            _replace_content_paragraph_preserving_style(row.cells[diary_col], diary_text)
            if day_col is not None:
                _set_simple_cell(row.cells[day_col], moment.strftime("%d"))
            if month_col is not None:
                _set_simple_cell(row.cells[month_col], moment.strftime("%m.%Y"))
            if hosp_col is not None:
                _set_simple_cell(row.cells[hosp_col], str(max(1, (moment.date() - admission.date()).days + 1)))
        else:
            _replace_content_paragraph_preserving_style(row.cells[diary_col], "")
            if day_col is not None:
                _set_simple_cell(row.cells[day_col], "")
            if month_col is not None:
                _set_simple_cell(row.cells[month_col], "")
            if hosp_col is not None:
                _set_simple_cell(row.cells[hosp_col], "")

    remove_forbidden_hospitalization_phrase_from_document(doc)
    output_root = Path(output_dir).expanduser()
    output_root.mkdir(parents=True, exist_ok=True)
    output = available_path(output_root / render_output_name(document, case, output_language=output_language, spellcheck_enabled=spellcheck_enabled))
    doc.save(str(output))

    try:
        from document_intelligence.form_fill import fill_docx_visible_fields

        fill_docx_visible_fields(
            output,
            context,
            role_id=document.role_id,
            category=document.category,
            button_label=document.button_label,
        )
    except Exception:
        try:
            output.unlink()
        except OSError:
            pass
        raise

    return CustomDiaryTemplateRenderResult(
        output,
        len(entries),
        sum(1 for _moment, _text, final in entries if final),
    )
