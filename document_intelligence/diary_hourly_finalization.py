"""Finalize text diaries when an intraday schedule is selected.

The calendar engine historically added the neutral discharge diary only in its
``daily`` branch.  The UI contract, however, exposes one ``force_final_diary``
choice independent of frequency.  This module closes that user-path gap without
re-enabling the removed legacy table writer.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from diary_dates import parse_optional_discharge_date
from diary_gender import adapt_text_to_patient_gender, detect_gender_from_patient_name
from diary_text_parser import remove_examinee_words
from diary_writer_apply import NEUTRAL_FINAL_DIARY_TEXT


def ensure_hourly_final_diary(
    result,
    *,
    discharge_value: str,
    patient_name: str,
    force_final_diary: bool,
) -> bool:
    """Append one neutral discharge diary when the hourly engine omitted it.

    Returns True only when a file was modified.  Existing final rows are never
    duplicated, and an invalid/missing discharge date is left to the normal
    preflight validation path.
    """

    if not force_final_diary or int(getattr(result, "final_rows_filled", 0) or 0) > 0:
        return False
    discharge = parse_optional_discharge_date(discharge_value)
    paths = list(getattr(result, "created_files", ()) or ())
    if discharge is None or not paths:
        return False

    target = Path(paths[0])
    if not target.exists() or target.suffix.lower() not in {".docx", ".docm"}:
        return False

    gender = detect_gender_from_patient_name(str(patient_name or ""))
    final_text, _changed = adapt_text_to_patient_gender(NEUTRAL_FINAL_DIARY_TEXT, gender)
    final_text = remove_examinee_words(final_text)

    doc = Document(str(target))
    if doc.paragraphs:
        doc.add_paragraph("")
    doc.add_paragraph(f"{discharge:%d.%m.%y} {final_text}".rstrip())
    from diary_batch import TEXT_DIARY_SIGNATURE_LINES

    for line in TEXT_DIARY_SIGNATURE_LINES:
        doc.add_paragraph(line)
    doc.save(str(target))

    # DiaryBatchResult is intentionally mutable; creation orchestration already
    # remaps its paths after transaction commit. Keep its accounting truthful.
    result.final_rows_filled = int(getattr(result, "final_rows_filled", 0) or 0) + 1
    result.filled_rows = int(getattr(result, "filled_rows", 0) or 0) + 1
    return True
