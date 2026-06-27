"""Regression smoke for the v1.4.11 auditor layer."""

from __future__ import annotations

import time
from pathlib import Path
import os

from docx import Document

from architecture_contracts import assert_architecture_contracts
from auditor_layer import audit_profile, audit_profile_and_case, audit_one_template, save_audit_report
from auditor_models import AuditSeverity, assert_auditor_models_lock
from auditor_profile import assert_auditor_profile_lock
from auditor_runtime import assert_auditor_runtime_lock
from auditor_template import assert_auditor_template_lock
from auditor_layer import assert_auditor_layer_lock
from universal_fields import PatientCase
from universal_profiles import DocumentPack, DocumentTemplateSpec

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "test_run_auditor_layer"
OUT.mkdir(exist_ok=True)


def _docx(path: Path, paragraphs: list[str]) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def run() -> None:
    assert_architecture_contracts()
    assert_auditor_models_lock()
    assert_auditor_template_lock()
    assert_auditor_profile_lock()
    assert_auditor_runtime_lock()
    assert_auditor_layer_lock()

    template = _docx(OUT / "operation_audit_template.docx", [
        "Протокол операции",
        "Пациент: {{patient.fio}}",
        "Номер истории болезни: {{case.number}}",
        "Операция: {{procedure.name}}",
        "Анестезия: {{procedure.anesthesia}}",
    ])
    template_report = audit_one_template(template)
    assert template_report.findings, "template audit must return structured findings"
    assert any(item.code == "LANGUAGE_DETECTED" and item.details.get("language_id") == "ru" for item in template_report.findings), template_report.human_report()

    inert = _docx(OUT / "inert_template.docx", ["Просто пустой бланк без placeholders"])
    inert_report = audit_one_template(inert)
    assert any(item.code == "TEMPLATE_HAS_NO_PLACEHOLDERS" and item.blocking for item in inert_report.findings), inert_report.human_report()

    pack = DocumentPack(
        pack_id="doctor.audit.test",
        name="Doctor Audit Test",
        specialty="surgery",
        documents=(
            DocumentTemplateSpec(
                id="operation_protocol_main",
                button_label="Протокол операции",
                template="templates/operation_audit_template.docx",
                required_fields=("patient.fio", "case.number", "procedure.name", "procedure.anesthesia"),
                role_id="operation_protocol",
            ),
        ),
    )
    templates_dir = OUT / "templates"
    templates_dir.mkdir(exist_ok=True)
    (templates_dir / template.name).write_bytes(template.read_bytes())
    profile_report = audit_profile(pack, base_dir=OUT)
    assert profile_report.ok, profile_report.human_report()

    case = PatientCase()
    case.update_from_pairs({"patient.fio": "Иванов Иван", "case.number": "42"}, source_document="smoke")
    runtime_report = audit_profile_and_case(pack, case, document_ids=("operation_protocol_main",), base_dir=OUT, output_dir=OUT)
    assert any(item.code == "GENERATION_HAS_MISSING_FIELDS" and not item.blocking for item in runtime_report.findings), runtime_report.human_report()
    report_path = save_audit_report(runtime_report, OUT / "auditor_report.txt")
    assert report_path.exists() and "Аудит" in report_path.read_text(encoding="utf-8")

    started = time.perf_counter()
    for _ in range(20):
        audit_one_template(template)
    elapsed = time.perf_counter() - started
    assert elapsed < 1.25, f"auditor cache/performance regressed: {elapsed:.3f}s"

    # The auditor layer must not be a hard blocker for doctor's own decision.
    assert not any(item.blocking for item in runtime_report.findings if item.code == "GENERATION_HAS_MISSING_FIELDS")
    print("AUDITOR LAYER SMOKE OK", flush=True)
    os._exit(0)


if __name__ == "__main__":
    run()
