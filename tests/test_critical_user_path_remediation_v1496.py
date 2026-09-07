from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from docx import Document

from custom_diary_template_renderer import render_custom_diary_template
from desktop_intake import (
    _is_supported_intake_document_name,
    _same_patient_existing_folder,
    normalize_intake_settings,
)
from desktop_intake_mixin import DesktopIntakeMixin
from diary_hourly_finalization import ensure_hourly_final_diary
from diary_models import DiaryBatchResult
from diary_schedule import DiaryScheduleSpec
from medical_models import PatientData
from medical_parser import MedicalTextParser
from universal_case_adapter import merge_case_values, patient_data_to_case
from universal_fields import PatientCase
from universal_profiles import DocumentPack, DocumentTemplateSpec


class _Var:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


def _case(**pairs: str) -> PatientCase:
    case = PatientCase()
    case.update_from_pairs(pairs, confidence=1.0, source_document="test")
    return case


def _make_diary_template(path: Path, *, rows: int = 6) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "ШАПКА ВРАЧА"
    doc.sections[0].footer.paragraphs[0].text = "ФУТЕР ВРАЧА"
    doc.add_paragraph("Пациент: {{patient.fio}}")
    table = doc.add_table(rows=rows + 1, cols=4)
    for index, value in enumerate(("Число", "Месяц", "День госпитализации", "Дневник наблюдения")):
        table.rows[0].cells[index].text = value
    for row in table.rows[1:]:
        diary_cell = row.cells[3]
        diary_cell.paragraphs[0].text = ""
        diary_cell.add_paragraph("Лечащий врач ____________________")
    doc.save(path)


def test_patient_case_does_not_invent_medical_semantics_and_derives_age() -> None:
    data = PatientData(
        fio="Иванов Иван Иванович",
        birth="01.01.1980",
        admission_date="02.01.2026",
        epi_text="Эпикриз не является результатом лечения",
        additional_info_text="Дополнительная информация не является рекомендациями",
        somatic_status="Соматический статус",
        profile_status="Психический статус",
        expert_work_status="да",
        sick_leave="12345",
    )
    case = patient_data_to_case(data)
    assert case.get("patient.age") == "46"
    assert case.get("anamnesis.expert") == ""
    assert case.get("condition.discharge") == ""
    assert case.get("treatment.result") == ""
    assert case.get("recommendations") == ""


def test_doctor_not_working_clears_stale_position_and_generic_recommendations() -> None:
    base = _case(**{"patient.position": "инженер", "expert.position": "инженер", "expert.work_org": "Завод"})
    merged = merge_case_values(
        base,
        {
            "patient.work": "не работает",
            "additional.info": "Сведения",
            "recommendations": "Сведения",
        },
        source_document="doctor",
    )
    assert merged.get("patient.work") == "не работает"
    assert merged.get("patient.position") == ""
    assert merged.get("expert.position") == ""
    assert merged.get("expert.work_org") == ""
    assert merged.get("recommendations") == ""
    assert merged.get("additional.info") == "Сведения"


def test_parser_never_treats_age_or_generic_treatment_as_birth_or_plan() -> None:
    parser = MedicalTextParser()
    assert "Возраст" not in parser.FIELD_ALIASES["birth"]
    assert "Wiek" not in parser.FIELD_ALIASES["birth"]
    assert "Лечение" not in parser.BLOCK_ALIASES["treatment_plan"]
    assert "Leczenie" not in parser.BLOCK_ALIASES["treatment_plan"]
    assert "Terapia" not in parser.BLOCK_ALIASES["treatment_plan"]


def test_custom_diary_preserves_doctor_word_template_and_final_row(tmp_path: Path) -> None:
    template = tmp_path / "doctor_diary.docx"
    _make_diary_template(template)
    case = _case(
        **{
            "patient.fio": "Иванов Иван Иванович",
            "admission.date": "01.01.2026",
            "discharge.date": "03.01.2026",
        }
    )
    spec = DocumentTemplateSpec(
        id="doctor_diary",
        button_label="Мои дневники",
        template=str(template),
        category="diaries",
        output_name="{{patient.fio}} {{document.label}}.docx",
    )
    result = render_custom_diary_template(
        template_path=template,
        output_dir=tmp_path / "out",
        case=case,
        document=spec,
        schedule=DiaryScheduleSpec("daily", (1, 2, 3), (), 1.0, "test"),
        statuses=("Состояние стабильное.", "Положительная динамика."),
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026",
        discharge_value="03.01.2026",
        repeat_statuses=True,
        force_final_diary=True,
    )
    assert result.path.exists()
    assert "Мои дневники" in result.path.name
    assert result.final_rows_filled == 1
    rendered = Document(result.path)
    assert rendered.sections[0].header.paragraphs[0].text == "ШАПКА ВРАЧА"
    assert rendered.sections[0].footer.paragraphs[0].text == "ФУТЕР ВРАЧА"
    assert "Пациент: Иванов Иван Иванович" in "\n".join(p.text for p in rendered.paragraphs)
    table_text = "\n".join(cell.text for row in rendered.tables[0].rows for cell in row.cells)
    assert "Состояние стабильное" in table_text
    assert "Лечащий врач" in table_text
    assert "03" in rendered.tables[0].rows[2].cells[0].text or "03" in table_text


def test_hourly_user_flow_gets_final_discharge_diary(tmp_path: Path) -> None:
    output = tmp_path / "hourly.docx"
    doc = Document()
    doc.add_paragraph("01.01.26 12:00 Наблюдение")
    doc.save(output)
    result = DiaryBatchResult([output], None, 1, 1, 1, 0, 0, 0, 0, 0)
    changed = ensure_hourly_final_diary(
        result,
        discharge_value="02.01.2026",
        patient_name="Иванов Иван Иванович",
        force_final_diary=True,
    )
    assert changed is True
    assert result.final_rows_filled == 1
    assert result.filled_rows == 2
    text = "\n".join(p.text for p in Document(output).paragraphs)
    assert "02.01.26" in text
    assert "Лечащий врач" in text


def test_custom_diary_set_rolls_back_when_any_selected_template_fails(tmp_path: Path, monkeypatch) -> None:
    import custom_diary_template_renderer
    import universal_diary_generation

    valid_template = tmp_path / "valid.docx"
    Document().save(valid_template)
    status_doc = tmp_path / "statuses.docx"
    Document().save(status_doc)
    pack = DocumentPack(
        pack_id="p",
        name="P",
        documents=(
            DocumentTemplateSpec("one", "Дневник один", str(valid_template), category="diaries"),
            DocumentTemplateSpec("two", "Дневник два", str(tmp_path / "missing.docx"), category="diaries"),
        ),
    )
    created = tmp_path / "created.docx"

    def fake_render(**_kwargs):
        created.write_bytes(b"created")
        return SimpleNamespace(path=created)

    monkeypatch.setattr(custom_diary_template_renderer, "render_custom_diary_template", fake_render)
    monkeypatch.setattr(universal_diary_generation, "read_statuses_from_files", lambda _paths: ["status"])
    monkeypatch.setattr(universal_diary_generation, "_effective_status_files", lambda _files, _template: (status_doc,))

    result = universal_diary_generation.render_diary_documents_from_pack(
        pack=pack,
        case=_case(**{"patient.fio": "Иванов И.И.", "admission.date": "01.01.2026", "discharge.date": "02.01.2026"}),
        document_ids=("one", "two"),
        output_dir=tmp_path,
        base_dir=None,
        status_files=(status_doc,),
        patient_name="Иванов И.И.",
        admission_value="01.01.2026",
        discharge_value="02.01.2026",
        diary_day_offsets=(1,),
    )
    assert result.created_files == ()
    assert result.skipped
    assert not created.exists()


def test_each_custom_diary_prefers_its_own_embedded_texts(monkeypatch, tmp_path: Path) -> None:
    import universal_diary_generation

    own = tmp_path / "own.docx"
    external = tmp_path / "external.docx"
    own.write_bytes(b"x")
    external.write_bytes(b"y")
    monkeypatch.setattr(
        universal_diary_generation,
        "extract_statuses_from_docx",
        lambda path: ["own status"] if Path(path) == own else [],
    )
    assert universal_diary_generation._effective_status_files((external,), own) == (own,)


def test_desktop_intake_supports_doc_and_exact_retry_without_fio(tmp_path: Path, monkeypatch) -> None:
    import desktop_patient_folder

    assert _is_supported_intake_document_name(Path("Первичный.doc"))
    source = tmp_path / "source.docx"
    source.write_bytes(b"same primary bytes")
    patient_dir = tmp_path / "patient"
    patient_dir.mkdir()
    existing = patient_dir / "old.docx"
    existing.write_bytes(b"same primary bytes")
    monkeypatch.setattr(
        desktop_patient_folder,
        "build_patient_folder_info",
        lambda _path: SimpleNamespace(fio="", admission_date=""),
    )
    assert _same_patient_existing_folder(patient_dir, source) is True


def test_intake_settings_keep_closed_app_agent_readiness() -> None:
    settings = normalize_intake_settings(
        {
            "asked": True,
            "enabled": True,
            "background_agent_ready": False,
            "folder": "Выписанные пациенты",
        }
    )
    assert settings["enabled"] is True
    assert settings["background_agent_ready"] is False


def test_intake_prompts_for_folder_naming_identity_before_strict_build(monkeypatch) -> None:
    import desktop_intake_mixin
    import desktop_patient_folder

    app = object.__new__(DesktopIntakeMixin)
    app.root = object()
    app.patient_name_var = _Var("")
    app.admission_date_var = _Var("")
    app.data = SimpleNamespace(admission_date="")
    app._set_ui_var = lambda var, value: var.set(value)
    monkeypatch.setattr(
        desktop_patient_folder,
        "build_patient_folder_info",
        lambda _path: SimpleNamespace(fio="", admission_date=""),
    )
    answers = iter(("Иванов Иван Иванович", "05.05.2026"))
    monkeypatch.setattr(desktop_intake_mixin.simpledialog, "askstring", lambda *_a, **_k: next(answers))

    fio, admission = app._prompt_intake_folder_identity(
        Path("primary.docx"),
        {"parts": ["full_fio", "admission_date"], "date_format": "short"},
    )
    assert fio == "Иванов Иван Иванович"
    assert admission == "05.05.2026"
    assert app.patient_name_var.get() == fio
    assert app.admission_date_var.get() == admission


def test_installer_does_not_create_intake_folder_before_user_consent() -> None:
    script = Path("installer/Dokkomplekt.iss").read_text(encoding="utf-8")
    assert "[Dirs]" not in script
    assert 'Name: "{userdesktop}\\Выписанные пациенты"' not in script
