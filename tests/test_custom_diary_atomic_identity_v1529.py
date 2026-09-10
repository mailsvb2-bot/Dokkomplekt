from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from diary_batch import fill_diary_batch
from universal_diary_generation import render_diary_documents_from_pack
from universal_fields import PatientCase
from universal_profiles import DocumentPack, DocumentTemplateSpec


def _case(*, discharge: str = "03.01.2026") -> PatientCase:
    case = PatientCase()
    case.update_from_pairs(
        {
            "patient.fio": "Иванов Иван Иванович",
            "admission.date": "01.01.2026 10:00",
            "discharge.date": discharge,
        },
        confidence=1.0,
        source_document="test",
    )
    return case


def _texts(tmp_path: Path) -> Path:
    path = tmp_path / "Тексты.docx"
    doc = Document()
    doc.add_paragraph("Состояние стабильное, контактен, назначения выполняет без замечаний.")
    doc.save(path)
    return path


def _template(tmp_path: Path, name: str) -> Path:
    path = tmp_path / name
    Document().save(path)
    return path


def test_selected_dates_without_in_stay_dates_are_rejected(tmp_path: Path) -> None:
    dates = tmp_path / "Даты.docx"
    doc = Document()
    doc.add_paragraph("20.01.2026")
    doc.save(dates)
    template = _template(tmp_path, "diary.docx")
    pack = DocumentPack(
        pack_id="invalid.dates",
        name="Invalid dates",
        documents=(
            DocumentTemplateSpec(
                id="diary",
                button_label="Дневники",
                template=str(template),
                category="diaries",
            ),
        ),
    )

    result = render_diary_documents_from_pack(
        pack=pack,
        case=_case(),
        document_ids=("diary",),
        output_dir=tmp_path / "out-invalid-dates",
        base_dir=None,
        status_files=(_texts(tmp_path),),
        date_files=(dates,),
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026 10:00",
        discharge_value="03.01.2026",
        diary_day_offsets=(1, 2),
    )

    assert not result.created_files
    assert result.skipped
    assert "не найдено подходящих дат" in result.skipped[0]
    assert not list((tmp_path / "out-invalid-dates").glob("*.docx"))


def test_equivalent_selected_diary_buttons_collapse_to_one_output(tmp_path: Path) -> None:
    first = _template(tmp_path, "first.docx")
    second = _template(tmp_path, "second.docx")
    pack = DocumentPack(
        pack_id="duplicate.buttons",
        name="Duplicate buttons",
        documents=(
            DocumentTemplateSpec(id="first", button_label="Дневники 1", template=str(first), category="diaries"),
            DocumentTemplateSpec(id="second", button_label="Дневники 2", template=str(second), category="diaries"),
        ),
    )

    result = render_diary_documents_from_pack(
        pack=pack,
        case=_case(),
        document_ids=("first", "second"),
        output_dir=tmp_path / "out-duplicate",
        base_dir=None,
        status_files=(_texts(tmp_path),),
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026 10:00",
        discharge_value="03.01.2026",
        diary_day_offsets=(1, 2),
    )

    assert not result.skipped
    assert len(result.created_files) == 1
    assert len(list((tmp_path / "out-duplicate").glob("*.docx"))) == 1
    assert any("создан один общий дневник" in warning for warning in result.warnings)


def test_report_failure_never_publishes_staged_custom_diary(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template = _template(tmp_path, "report-diary.docx")
    pack = DocumentPack(
        pack_id="report.failure",
        name="Report failure",
        documents=(DocumentTemplateSpec(id="diary", button_label="Дневники", template=str(template), category="diaries"),),
    )
    output_dir = tmp_path / "out-report-failure"

    def fail_report_path(*_args, **_kwargs):
        raise OSError("report boom")

    monkeypatch.setattr("diary_batch.technical_report_path", fail_report_path)

    result = render_diary_documents_from_pack(
        pack=pack,
        case=_case(),
        document_ids=("diary",),
        output_dir=output_dir,
        base_dir=None,
        status_files=(_texts(tmp_path),),
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026",
        discharge_value="03.01.2026",
        diary_day_offsets=(1, 2),
        write_report=True,
    )

    assert not result.created_files
    assert result.skipped
    assert "report boom" in result.skipped[0]
    assert output_dir.exists()
    assert not list(output_dir.iterdir())
