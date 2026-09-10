from __future__ import annotations

from pathlib import Path

from docx import Document

from universal_diary_generation import render_diary_documents_from_pack
from universal_fields import PatientCase
from universal_profiles import DocumentPack, DocumentTemplateSpec


def _text(path: Path) -> str:
    return "\n".join(p.text for p in Document(str(path)).paragraphs if p.text.strip())


def _case(discharge: str) -> PatientCase:
    case = PatientCase()
    case.update_from_pairs(
        {
            "patient.fio": "Иванов Иван Иванович",
            "admission.date": "01.01.2026 10:00",
            "discharge.date": discharge,
        },
        confidence=1.0,
        source_document="test",
    )
    return case


def _pack(template: Path, doc_id: str) -> DocumentPack:
    return DocumentPack(
        pack_id=doc_id,
        name=doc_id,
        documents=(
            DocumentTemplateSpec(
                id=doc_id,
                button_label="Почасовые дневники",
                template=str(template),
                category="diaries",
            ),
        ),
    )


def _texts_file(tmp_path: Path) -> Path:
    path = tmp_path / "Тексты.docx"
    doc = Document()
    doc.add_paragraph("Состояние стабильное, контактен, назначения выполняет без замечаний.")
    doc.save(path)
    return path


def test_selected_dates_constrain_hourly_custom_diary(tmp_path: Path) -> None:
    template = tmp_path / "hourly.docx"
    Document().save(template)
    dates = tmp_path / "Даты.docx"
    doc = Document()
    doc.add_paragraph("07.01.2026")
    doc.save(dates)

    result = render_diary_documents_from_pack(
        pack=_pack(template, "hourly_dates"),
        case=_case("08.01.2026"),
        document_ids=("hourly_dates",),
        output_dir=tmp_path / "out-dates",
        base_dir=None,
        status_files=(_texts_file(tmp_path),),
        date_files=(dates,),
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026 10:00",
        discharge_value="08.01.2026",
        frequency_mode="hourly",
        diary_hour_offsets=(24,),
        force_final_diary=False,
    )

    assert not result.skipped
    text = _text(Path(result.created_files[0]))
    assert "07.01.26 10:00" in text
    assert "08.01.26 10:00" in text
    assert "02.01.26" not in text
    assert "03.01.26" not in text


def test_hourly_custom_diary_appends_one_discharge_entry(tmp_path: Path) -> None:
    template = tmp_path / "hourly-final.docx"
    Document().save(template)
    result = render_diary_documents_from_pack(
        pack=_pack(template, "hourly_final"),
        case=_case("03.01.2026"),
        document_ids=("hourly_final",),
        output_dir=tmp_path / "out-final",
        base_dir=None,
        status_files=(_texts_file(tmp_path),),
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026 10:00",
        discharge_value="03.01.2026",
        frequency_mode="hourly",
        diary_hour_offsets=(24,),
        force_final_diary=True,
    )

    assert not result.skipped
    text = _text(Path(result.created_files[0]))
    marker = "03.01.26 Состояние улучшилось"
    assert marker in text
    assert text.count(marker) == 1
