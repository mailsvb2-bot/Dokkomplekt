from __future__ import annotations

from pathlib import Path

from actions_universal_flow import ActionsUniversalFlowMixin
from diary_creation_wizard import build_diary_wizard_review, confirm_diary_creation


class _Var:
    def __init__(self, value: str):
        self._value = value

    def get(self) -> str:
        return self._value


class _Document:
    name = "Дневн"
    template = "templates/doctor_diary.docx"


class _Pack:
    def document_by_id(self, document_id: str):
        return _Document() if document_id == "diary" else None


class _App(ActionsUniversalFlowMixin):
    root = None
    status_files: list[str] = []
    _diary_text_files_auto_selected = False
    _last_diary_wizard_review = None

    patient_name_var = _Var("Тестовый Пациент")
    admission_date_var = _Var("30.09.2025")
    discharge_date_var = _Var("05.06.2026")
    diary_frequency_mode_var = _Var("daily")
    expert_sick_leave_needed_var = _Var("нет")

    def _universal_profile_path(self) -> Path:
        return Path("/tmp/profile/medpack.json")

    def _auto_select_diary_text_by_diagnosis(self, *, ask_folder: bool = False) -> bool:
        return False

    def choose_status_files(self) -> None:
        return None


def test_embedded_doctor_template_is_accepted_as_diary_text_source(monkeypatch):
    import universal_diary_generation

    monkeypatch.setattr(
        universal_diary_generation,
        "diary_documents_have_embedded_status_texts",
        lambda **_kwargs: True,
    )
    app = _App()

    labels = app._ensure_diary_text_files_for_creation(_Pack(), ["diary"])

    assert labels == ("тексты внутри шаблона: Дневн",)
    assert app.status_files == []
    assert confirm_diary_creation(app, text_source_labels=labels) is True
    assert app._last_diary_wizard_review is not None
    assert app._last_diary_wizard_review.ok
    assert app._last_diary_wizard_review.text_files == labels


def test_wizard_still_rejects_missing_text_source_without_override():
    app = _App()
    review = build_diary_wizard_review(app)

    assert not review.ok
    assert "Выберите тексты дневников." in review.warnings


def test_embedded_template_texts_flow_through_real_renderer(tmp_path: Path):
    from docx import Document
    from universal_diary_generation import render_diary_documents_from_pack
    from universal_fields import PatientCase
    from universal_profiles import DocumentPack, DocumentTemplateSpec

    template = tmp_path / "doctor_diary.docx"
    doc = Document()
    doc.add_paragraph("Состояние стабильное, контактен, назначения выполняет, режим соблюдает.")
    doc.add_paragraph("Жалоб активно не предъявляет, отрицательной динамики не отмечается.")
    doc.save(template)

    case = PatientCase()
    case.update_from_pairs(
        {
            "patient.fio": "Тестовый Пациент",
            "admission.date": "30.09.2025",
            "discharge.date": "05.06.2026",
        },
        confidence=1.0,
        source_document="screenshot-regression",
    )
    pack = DocumentPack(
        pack_id="embedded.diary",
        name="Embedded diary",
        documents=(
            DocumentTemplateSpec(
                id="diary",
                button_label="Дневн",
                template=str(template),
                category="diaries",
                role_id="daily_diary",
            ),
        ),
    )

    result = render_diary_documents_from_pack(
        pack=pack,
        case=case,
        document_ids=("diary",),
        output_dir=tmp_path / "out",
        base_dir=None,
        status_files=(),
        patient_name="Тестовый Пациент",
        admission_value="30.09.2025",
        discharge_value="05.06.2026",
        diary_day_offsets=(1, 2),
        force_final_diary=True,
    )

    assert not result.skipped
    assert len(result.created_files) == 1
    rendered = Document(str(result.created_files[0]))
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "Состояние стабильное" in text
    assert "Жалоб активно не предъявляет" in text
    assert "Состояние улучшилось" in text
