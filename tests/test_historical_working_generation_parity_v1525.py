from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from actions_creation_orchestrator import ActionsCreationLiveGuardMixin
from universal_fields import PatientCase
from universal_generation import render_documents_from_pack
from universal_profiles import DocumentPack
from universal_template_engine import infer_document_spec_from_template


def _template(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Пациент: {{patient.fio}}")
    doc.add_paragraph("Диагноз: {{diagnosis.main}}")
    # Real doctor profiles may persist this explicit placeholder as optional.
    # Its absence must not cancel the entire document transaction.
    doc.add_paragraph("Номер больничного: {{expert.sick_leave_number}}")
    doc.save(path)
    return path


def _pack(tmp_path: Path) -> DocumentPack:
    template = _template(tmp_path / "doctor.docx")
    inferred = infer_document_spec_from_template(
        template,
        button_label="Документ врача",
        document_id="doctor_document",
        role_id="discharge",
    )
    spec = replace(
        inferred,
        template=template.name,
        required_fields=("patient.fio", "diagnosis.main"),
        optional_fields=("expert.sick_leave_number",),
    )
    return DocumentPack(pack_id="historical.compat", name="Historical compat", documents=(spec,))


class _ActualRendererParent:
    def _create_regular_custom_documents(self, current_pack, case, regular_ids, out_dir):
        result = render_documents_from_pack(
            pack=current_pack,
            case=case,
            document_ids=regular_ids,
            output_dir=out_dir,
            base_dir=Path(out_dir).parent,
            strict=not bool(getattr(self, "_allow_missing_required_creation", False)),
        )
        if result.skipped_documents:
            raise ValueError("; ".join(result.skipped_documents))
        return [Path(item) for item in result.created_files]


class _Subject(ActionsCreationLiveGuardMixin, _ActualRendererParent):
    pass


def test_working_historical_rule_allows_empty_optional_placeholder(tmp_path: Path) -> None:
    pack = _pack(tmp_path)
    case = PatientCase()
    case.set("patient.fio", "Орлова Мария Ивановна")
    case.set("diagnosis.main", "F32.1 Депрессивный эпизод")
    app = _Subject()
    app._allow_missing_required_creation = False

    created = app._create_regular_custom_documents(pack, case, ["doctor_document"], tmp_path / "out")

    assert len(created) == 1
    assert created[0].exists()
    text = "\n".join(paragraph.text for paragraph in Document(created[0]).paragraphs)
    assert "Орлова Мария Ивановна" in text
    assert "F32.1 Депрессивный эпизод" in text
    assert "{{expert.sick_leave_number}}" not in text
    assert app._allow_missing_required_creation is False


def test_historical_compatibility_does_not_soften_required_fields(tmp_path: Path) -> None:
    pack = _pack(tmp_path)
    case = PatientCase()
    case.set("patient.fio", "Орлова Мария Ивановна")
    app = _Subject()
    app._allow_missing_required_creation = False

    with pytest.raises(ValueError, match="diagnosis.main"):
        app._create_regular_custom_documents(pack, case, ["doctor_document"], tmp_path / "out")

    assert not (tmp_path / "out").exists()
    assert app._allow_missing_required_creation is False


def test_real_app_mro_uses_historical_compatibility_boundary() -> None:
    from app import CombinedMedicalDiaryApp

    assert (
        CombinedMedicalDiaryApp._create_regular_custom_documents
        is ActionsCreationLiveGuardMixin._create_regular_custom_documents
    )
