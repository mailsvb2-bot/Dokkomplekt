from __future__ import annotations

from pathlib import Path

from docx import Document

from medical_docx_reader import extract_docx_text
from medical_parser import MedicalTextParser
from universal_case_adapter import patient_data_to_case
from universal_generation import render_documents_from_pack
from universal_profiles import DocumentPack
from universal_template_engine import infer_document_spec_from_template


VALID_DIAGNOSIS = "F20.00 Параноидная шизофрения"


def _build_primary(path: Path) -> Path:
    """Create a realistic primary Word document with data in body and a table."""

    document = Document()
    document.add_paragraph("01.09.2026 Первичный осмотр")
    document.add_paragraph("История болезни № 321")
    document.add_paragraph("Ф.И.О.: Петров Петр Петрович")
    document.add_paragraph("Дата рождения: 15.04.1985")
    document.add_paragraph("Дата поступления: 01.09.2026")
    document.add_paragraph("Дата выписки: 08.09.2026")
    document.add_paragraph("Лечение: рисперидон 4 мг/сут, наблюдение врача")
    document.add_paragraph(f"Диагноз: {VALID_DIAGNOSIS}")
    work = document.add_table(rows=2, cols=1)
    work.cell(0, 0).text = "Место работы: АО Пример"
    work.cell(1, 0).text = "Должность: инженер"
    document.save(path)
    return path


def _build_doctor_template(path: Path) -> Path:
    """Create one doctor-owned template spanning body, table, header and footer."""

    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Пациент {{patient.fio}}"
    section.footer.paragraphs[0].text = "История болезни № {{case.number}}"

    document.add_paragraph("ФИО: {{patient.fio}}")
    document.add_paragraph("Дата рождения: {{patient.birth_date}}")
    document.add_paragraph("Дата поступления: {{admission.date}}")
    document.add_paragraph("Дата выписки: {{discharge.date}}")
    document.add_paragraph("Диагноз: {{diagnosis.main}}")
    document.add_paragraph("Лечение: {{treatment.plan}}")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Место работы"
    table.cell(0, 1).text = "{{patient.work}}"
    table.cell(1, 0).text = "Должность"
    table.cell(1, 1).text = "{{patient.position}}"
    document.save(path)
    return path


def test_primary_docx_data_reaches_real_filled_doctor_docx(tmp_path: Path) -> None:
    """Prove the real data chain: primary DOCX -> parser -> case -> filled DOCX."""

    primary = _build_primary(tmp_path / "Первичный осмотр.docx")
    data = MedicalTextParser().parse_docx(primary)

    # These values must come from the primary document itself, without popup or
    # manually injected PatientData.  Work/position additionally prove that the
    # primary reader does not ignore Word table content.
    assert data.fio == "Петров Петр Петрович"
    assert data.case_number == "321"
    assert data.birth == "15.04.1985"
    assert data.admission_date == "01.09.2026"
    assert data.discharge_date == "08.09.2026"
    assert data.diagnosis == VALID_DIAGNOSIS
    assert data.treatment_plan == "рисперидон 4 мг/сут, наблюдение врача"
    assert data.work_org == "АО Пример"
    assert data.position == "инженер"

    case = patient_data_to_case(data, source_document=str(primary))
    assert case.get("patient.fio") == "Петров Петр Петрович"
    assert case.get("case.number") == "321"
    assert case.get("patient.birth_date") == "15.04.1985"
    assert case.get("admission.date") == "01.09.2026"
    assert case.get("discharge.date") == "08.09.2026"
    assert case.get("diagnosis.main") == VALID_DIAGNOSIS
    assert case.get("treatment.plan") == "рисперидон 4 мг/сут, наблюдение врача"
    assert case.get("patient.work") == "АО Пример"
    assert case.get("patient.position") == "инженер"

    template = _build_doctor_template(tmp_path / "Шаблон врача.docx")
    spec = infer_document_spec_from_template(
        template,
        button_label="Контрольный документ",
        document_id="primary_data_proof",
        role_id="primary_exam",
    )
    pack = DocumentPack(pack_id="primary.data.proof", name="Primary data proof", documents=(spec,))

    result = render_documents_from_pack(
        pack=pack,
        case=case,
        document_ids=(spec.id,),
        output_dir=tmp_path / "result",
        base_dir=tmp_path,
        strict=True,
        output_language="ru",
        spellcheck_enabled=False,
    )

    assert result.ok, result.human_report()
    assert len(result.created_files) == 1
    output = Path(result.created_files[0])
    assert output.exists()

    # Re-open the physically saved DOCX.  This is the final user-visible file,
    # not the in-memory context.  extract_docx_text includes body, tables,
    # headers and footers, so it also locks template-fidelity paths.
    visible = extract_docx_text(output)
    for expected in (
        "Петров Петр Петрович",
        "321",
        "15.04.1985",
        "01.09.2026",
        "08.09.2026",
        VALID_DIAGNOSIS,
        "рисперидон 4 мг/сут, наблюдение врача",
        "АО Пример",
        "инженер",
    ):
        assert expected in visible, (expected, visible)
    assert "{{" not in visible and "}}" not in visible
