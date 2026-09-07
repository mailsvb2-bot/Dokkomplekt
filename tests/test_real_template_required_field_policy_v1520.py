from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document

from universal_fields import PatientCase
from universal_profiles import DocumentPack
from universal_template_engine import (
    infer_document_spec_from_template,
    infer_template_semantic_fields,
    migrate_auto_inferred_required_fields,
    missing_required_fields,
    render_template_to_docx,
)


def _realistic_discharge_template(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Выписной эпикриз")
    for text in (
        "ФИО: __________",
        "№ истории болезни: __________",
        "Диагноз: __________",
        "Лечение: __________",
        "Дата выписки: __________",
        "Номер больничного: __________",
        "Место работы: __________",
        "Должность: __________",
    ):
        document.add_paragraph(text)
    document.save(path)
    return path


def _core_case() -> PatientCase:
    case = PatientCase()
    case.set("patient.fio", "Тестовый Пациент")
    case.set("case.number", "123")
    case.set("diagnosis.main", "F20.0")
    case.set("treatment.plan", "Назначенная терапия")
    case.set("discharge.date", "07.09.2026")
    return case


def test_conditional_visible_fields_do_not_block_strict_real_template_render(tmp_path: Path) -> None:
    template = _realistic_discharge_template(tmp_path / "discharge.docx")
    spec = infer_document_spec_from_template(
        template,
        button_label="Выписной эпикриз",
        role_id="discharge_epicrisis",
    )

    assert spec.required_fields == (
        "patient.fio",
        "case.number",
        "diagnosis.main",
        "treatment.plan",
        "discharge.date",
    )
    assert {"expert.sick_leave_number", "patient.work", "patient.position"}.issubset(spec.optional_fields)
    assert missing_required_fields(_core_case(), spec) == ()

    output = tmp_path / "created.docx"
    render_template_to_docx(
        template_path=template,
        output_path=output,
        case=_core_case(),
        document=spec,
        strict=True,
    )
    assert output.exists()


def test_existing_auto_inferred_profile_is_migrated_without_recreating_buttons(tmp_path: Path) -> None:
    template = _realistic_discharge_template(tmp_path / "discharge.docx")
    fresh = infer_document_spec_from_template(
        template,
        button_label="Выписной эпикриз",
        role_id="discharge_epicrisis",
    )
    semantic_fields = infer_template_semantic_fields(
        template,
        role_id=fresh.role_id,
        category=fresh.category,
        button_label=fresh.button_label,
    )
    old_broken = replace(fresh, required_fields=semantic_fields, optional_fields=())
    pack = DocumentPack(pack_id="doctor.test", name="Test", documents=(old_broken,))

    assert migrate_auto_inferred_required_fields(pack, base_dir=tmp_path) is True
    migrated = pack.documents[0]
    assert migrated.id == old_broken.id
    assert migrated.template == old_broken.template
    assert migrated.button_label == old_broken.button_label
    assert migrated.required_fields == fresh.required_fields
    assert migrated.optional_fields == fresh.optional_fields


def test_explicit_required_field_policy_is_not_softened_by_migration(tmp_path: Path) -> None:
    template = _realistic_discharge_template(tmp_path / "discharge.docx")
    fresh = infer_document_spec_from_template(
        template,
        button_label="Выписной эпикриз",
        role_id="discharge_epicrisis",
    )
    explicit = replace(
        fresh,
        required_fields=(*fresh.required_fields, "expert.sick_leave_number"),
        optional_fields=tuple(field for field in fresh.optional_fields if field != "expert.sick_leave_number"),
        description="Политика обязательных полей задана врачом вручную.",
    )
    pack = DocumentPack(pack_id="doctor.explicit", name="Explicit", documents=(explicit,))

    assert migrate_auto_inferred_required_fields(pack, base_dir=tmp_path) is False
    assert pack.documents[0] == explicit


def test_explicit_placeholder_remains_required_even_when_not_a_role_core_field(tmp_path: Path) -> None:
    template = tmp_path / "explicit.docx"
    document = Document()
    document.add_paragraph("Документ {{patient.fio}}")
    document.add_paragraph("Дополнительное обязательное значение: {{custom.required_note}}")
    document.save(template)

    from universal_fields import FieldDefinition, default_field_registry

    registry = default_field_registry((FieldDefinition("custom.required_note", "Обязательная заметка", "custom"),))
    spec = infer_document_spec_from_template(
        template,
        button_label="Документ",
        role_id="unknown",
        registry=registry,
    )
    assert "custom.required_note" in spec.required_fields
    assert "custom.required_note" not in spec.optional_fields
