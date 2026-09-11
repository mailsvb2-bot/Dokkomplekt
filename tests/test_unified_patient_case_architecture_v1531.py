from pathlib import Path

from docx import Document

from medical_models import PatientData, build_patient_case_review
from universal_case_adapter import patient_data_to_case, supplement_patient_case
from universal_fields import PatientCase
from universal_profiles import DocumentTemplateSpec
from universal_template_engine import render_template_to_docx


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def test_output_identity_never_becomes_medical_fio():
    data = PatientData(fio="", output_fio="Сидорова С.С.", admission_date="01.09.2026")
    case = patient_data_to_case(data, source_document="primary.docx")
    assert case.get("patient.fio") == ""
    review = build_patient_case_review(data, selected_medical=("discharge",))
    assert review.value("output_fio") == "Сидорова С.С."
    assert [field.key for field in review.critical_missing()] == ["fio", "case_number", "discharge_date", "diagnosis", "treatment"]


def test_profile_status_keeps_neutral_and_psychiatric_semantic_keys():
    data = PatientData(profile_status="Спокоен, ориентирован, продуктивному контакту доступен.")
    case = patient_data_to_case(data, source_document="primary.docx")
    assert case.get("status.specialty") == data.profile_status
    assert case.get("status.mental") == data.profile_status


def test_profile_scanner_can_only_supplement_canonical_case():
    base = PatientCase()
    base.set("patient.fio", "Петров Пётр Петрович", confidence=0.60, source_document="primary")
    base.set("diagnosis.main", "F20.00 Шизофрения", confidence=0.60, source_document="primary")
    scan = PatientCase()
    scan.set("patient.fio", "Иванов Иван Иванович", confidence=1.0, source_document="profile_scan")
    scan.set("diagnosis.main", "F41.2 Старый диагноз", confidence=1.0, source_document="profile_scan")
    scan.set("custom.specialty_note", "Дополнение", confidence=1.0, source_document="profile_scan")
    merged = supplement_patient_case(base, scan)
    assert merged.get("patient.fio") == "Петров Пётр Петрович"
    assert merged.get("diagnosis.main") == "F20.00 Шизофрения"
    assert merged.get("custom.specialty_note") == "Дополнение"


def test_prefilled_doctor_template_is_layout_not_patient_database(tmp_path: Path):
    template = tmp_path / "doctor_discharge.docx"
    output = tmp_path / "rendered.docx"
    doc = Document()
    doc.add_paragraph("16.09.2026 Выписной эпикриз № 1023/з")
    doc.add_paragraph("Иванов Иван Иванович, 04.01.200 г.р., зарегистрирован по адресу: Старый адрес")
    doc.add_paragraph("Находился на лечении в ГБУЗ НО «НКЦПЗ» диспансер №2 с 12.07.2026 по 16.09.2026")
    doc.add_paragraph("Анамнез заболевания: ТЕХНИЧЕСКИЙ ТЕКСТ ИЗ СТАРОГО ШАБЛОНА")
    doc.add_paragraph("Психический статус при поступлении: СТАРЫЙ ПСИХИЧЕСКИЙ СТАТУС")
    doc.add_paragraph("Диагноз: F41.2 Старый диагноз")
    doc.add_paragraph("Лечение: Тералиджен 5 мг на ночь")
    doc.add_paragraph("ЭПИ – сюда подставлять информацию из файла ЭПИ")
    doc.save(template)

    case = PatientCase()
    case.update_from_pairs(
        {
            "patient.fio": "Петров Пётр Петрович",
            "patient.birth_date": "05.02.1990",
            "patient.address": "Нижний Новгород, Новая улица 1",
            "case.number": "77/к",
            "admission.date": "30.09.2026",
            "discharge.date": "09.10.2026",
            "anamnesis.disease": "Канонический анамнез текущего пациента.",
            "status.mental": "Канонический психический статус текущего пациента.",
            "diagnosis.main": "F20.00 Шизофрения параноидная",
            "treatment.plan": "Каноническое лечение текущего пациента.",
            "epicrisis.text": "Канонический текст ЭПИ текущего пациента.",
        },
        confidence=1.0,
        source_document="primary.docx",
    )
    spec = DocumentTemplateSpec(
        id="discharge",
        button_label="Выписной эпикриз",
        template=str(template),
        role_id="discharge_epicrisis",
        required_fields=("patient.fio", "admission.date", "discharge.date", "diagnosis.main", "treatment.plan"),
        optional_fields=("patient.birth_date", "patient.address", "anamnesis.disease", "status.mental", "epicrisis.text"),
    )
    result = render_template_to_docx(template_path=template, output_path=output, case=case, document=spec, strict=True)
    assert Path(result.output_path).exists()
    text = _docx_text(output)
    assert "Петров Пётр Петрович, 05.02.1990 г.р., зарегистрирован по адресу: Нижний Новгород, Новая улица 1" in text
    assert "09.10.2026 Выписной эпикриз № 77/к" in text
    assert "с 30.09.2026 по 09.10.2026" in text
    assert "Канонический анамнез текущего пациента." in text
    assert "Канонический психический статус текущего пациента." in text
    assert "F20.00 Шизофрения параноидная" in text
    assert "Каноническое лечение текущего пациента." in text
    assert "Канонический текст ЭПИ текущего пациента." in text
    for stale in ("Иванов Иван Иванович", "04.01.200", "F41.2", "ТЕХНИЧЕСКИЙ ТЕКСТ", "СТАРЫЙ ПСИХИЧЕСКИЙ", "Тералиджен 5 мг"):
        assert stale not in text


def test_explicit_placeholders_are_not_reinterpreted_by_visible_semantic_pass(tmp_path: Path):
    template = tmp_path / "combined_dates.docx"
    output = tmp_path / "combined_dates_rendered.docx"
    doc = Document()
    doc.add_paragraph("ФИО: {{patient.fio}}")
    doc.add_paragraph("Поступил: {{admission.date}}  Выписан: {{discharge.date}}")
    doc.add_paragraph("Диагноз: {{diagnosis.main}}")
    doc.save(template)

    case = PatientCase()
    case.update_from_pairs(
        {
            "patient.fio": "Орлова Мария Ивановна",
            "admission.date": "05.05.2026",
            "discharge.date": "19.05.2026",
            "diagnosis.main": "F32.1 Депрессивный эпизод средней степени",
        },
        source_document="primary.docx",
    )
    spec = DocumentTemplateSpec(
        id="discharge",
        button_label="Выписной эпикриз",
        template=str(template),
        role_id="discharge_epicrisis",
        required_fields=("patient.fio", "admission.date", "discharge.date", "diagnosis.main"),
    )

    render_template_to_docx(template_path=template, output_path=output, case=case, document=spec, strict=True)
    text = _docx_text(output)
    assert "Поступил: 05.05.2026" in text
    assert "Выписан: 19.05.2026" in text
    assert "{{" not in text


def test_consistency_gate_deletes_document_that_cannot_match_canonical_identity(tmp_path: Path):
    template = tmp_path / "unmapped_medical_template.docx"
    output = tmp_path / "must_not_survive.docx"
    doc = Document()
    doc.add_paragraph("Служебная форма без поля пациента")
    doc.add_paragraph("Чужой пациент остаётся только как нерегистрируемый текст")
    doc.save(template)

    case = PatientCase()
    case.set("patient.fio", "Смирнова Анна Викторовна", source_document="primary.docx")
    spec = DocumentTemplateSpec(
        id="primary_exam",
        button_label="Первичный осмотр",
        template=str(template),
        role_id="primary_exam",
    )

    import pytest
    with pytest.raises(ValueError, match="канонической карточкой пациента"):
        render_template_to_docx(
            template_path=template,
            output_path=output,
            case=case,
            document=spec,
            strict=True,
        )
    assert not output.exists()
