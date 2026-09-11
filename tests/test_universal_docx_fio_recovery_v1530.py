from __future__ import annotations

from docx import Document

from desktop_patient_folder import build_patient_folder_info, build_patient_folder_name
from medical_parser import MedicalTextParser


def test_invisible_word_separator_does_not_hide_unlabeled_fio():
    data = MedicalTextParser().parse_text(
        "15.04.2026 Направление на госпитализацию\n"
        "Иванов\ufffeИван Иванович, 04.01.2000, Нижний Новгород\n"
        "Диагноз: F20.00 Шизофрения"
    )
    assert data.fio == "Иванов Иван Иванович"
    assert build_patient_folder_name(
        fio=data.fio,
        settings={"parts": ["surname_initials"], "date_format": "short"},
        strict=True,
    ) == "Иванов И.И"


def test_invisible_word_formatting_is_removed_inside_labeled_fio():
    data = MedicalTextParser().parse_text(
        "15.04.2026 Первичный осмотр\n"
        "Ф.И.О.: Петрова\u200b Мария\ufeff Сергеевна\n"
        "Дата рождения: 04.01.2000"
    )
    assert data.fio == "Петрова Мария Сергеевна"


def test_table_header_row_recovers_any_patient_fio_for_strict_folder(tmp_path):
    primary = tmp_path / "primary.docx"
    doc = Document()
    doc.add_paragraph("15.04.2026 Направление на госпитализацию")
    table = doc.add_table(rows=2, cols=5)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Ф.И.О. больного"
    table.cell(0, 3).text = "Дата рождения"
    table.cell(0, 4).text = "Адрес регистрации"
    table.cell(1, 0).text = "Орлов"
    table.cell(1, 1).text = "Пётр"
    table.cell(1, 2).text = "Максимович"
    table.cell(1, 3).text = "04.01.2000"
    table.cell(1, 4).text = "Нижний Новгород"
    doc.add_paragraph("Диагноз: F20.00 Шизофрения")
    doc.save(primary)

    data = MedicalTextParser().parse_docx(primary)
    assert data.fio == "Орлов Пётр Максимович"
    assert build_patient_folder_info(primary).fio == "Орлов Пётр Максимович"
    assert build_patient_folder_name(
        fio=data.fio,
        settings={"parts": ["surname_initials"], "date_format": "short"},
        strict=True,
    ) == "Орлов П.М"


def test_table_geometry_never_promotes_clinician_fio_to_patient(tmp_path):
    primary = tmp_path / "clinician_only.docx"
    doc = Document()
    doc.add_paragraph("15.04.2026 Первичный осмотр")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Ф.И.О. лечащего врача"
    table.cell(0, 1).text = "Дата рождения"
    table.cell(1, 0).text = "Соколова Марина Андреевна"
    table.cell(1, 1).text = "04.01.2000"
    doc.add_paragraph("Диагноз: F20.00 Шизофрения")
    doc.save(primary)

    assert MedicalTextParser().parse_docx(primary).fio == ""
