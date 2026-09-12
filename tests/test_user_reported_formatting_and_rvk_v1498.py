from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

from diary_batch import fill_diary_batch
from medical_models import PatientData
from medical_renderer import MedicalDocumentRenderer
from universal_fields import PatientCase
from universal_profiles import DocumentTemplateSpec
from universal_template_engine import render_template_to_docx


class _Var:
    def __init__(self, value: str = "") -> None:
        self.value = value

    def get(self) -> str:
        return self.value

    def set(self, value: str) -> None:
        self.value = value


def _status_docx(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Состояние стабильное, назначения выполняет, жалоб активно не предъявляет.")
    doc.save(path)
    return path


def test_text_diary_uses_source_doctor_and_head_and_right_aligns_both(tmp_path: Path) -> None:
    source = _status_docx(tmp_path / "F20.0 шизофрения.docx")
    result = fill_diary_batch(
        status_files=(source,), diary_files=(), output_dir=tmp_path / "out",
        patient_name="Банина Екатерина Сергеевна", admission_value="01.09.2026",
        diary_day_offsets=(1,), force_final_diary=False,
        treating_physician="Лечащий врач Балаганин С.В.",
        department_head="Зав. отделением Можарова Е.А.",
    )
    doc = Document(str(result.created_files[0]))
    signatures = [p for p in doc.paragraphs if p.text.startswith(("Лечащий врач", "Зав.отделением"))]
    assert signatures
    assert len(signatures) % 2 == 0
    assert [p.text for p in signatures] == [
        expected
        for _ in range(len(signatures) // 2)
        for expected in (
            "Лечащий врач Балаганин С.В.",
            "Зав.отделением Можарова Е.А.",
        )
    ]
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.RIGHT for p in signatures)


def _discharge_template(path: Path) -> Path:
    doc = Document()
    red = doc.styles.add_style("Red treatment period", WD_STYLE_TYPE.PARAGRAPH)
    red.font.color.rgb = RGBColor(255, 0, 0)
    doc.add_paragraph("Дата, время")
    doc.add_paragraph("Ф.И.О., г.р., зарегистрирован по адресу")
    period = doc.add_paragraph("Находился на лечении в старом учреждении с 01.01.2000 по 02.01.2000")
    period.style = red
    doc.add_paragraph("В 3 отделение КДП поступает")
    doc.add_paragraph("Жалобы при поступлении")
    doc.add_paragraph("Анамнез жизни")
    doc.add_paragraph("Анамнез заболевания")
    doc.add_paragraph("Сомато-неврологический статус")
    doc.add_paragraph("Зав. отд. __________ Врач __________")
    doc.save(path)
    return path


def test_discharge_period_stays_black_after_female_gender_pass_and_mode_is_rendered(tmp_path: Path) -> None:
    template = _discharge_template(tmp_path / "discharge.docx")
    output = tmp_path / "out.docx"
    data = PatientData(
        fio="Банина Екатерина Сергеевна",
        admission_date="01.09.2026",
        discharge_date="09.09.2026",
        admission_mode="повторно",
        doctor="Балаганин С.В.",
        head="Можарова Е.А.",
    )
    MedicalDocumentRenderer().render("discharge", template, output, data)
    doc = Document(str(output))
    period = next(p for p in doc.paragraphs if p.text.startswith("Находилась на лечении"))
    assert period.runs
    assert all(run.font.color.rgb == RGBColor(0, 0, 0) for run in period.runs)
    assert any(p.text == "В 3 отделение КДП поступает повторно" for p in doc.paragraphs)


def test_rvk_builtin_renders_primary_admission_mode(tmp_path: Path) -> None:
    template = tmp_path / "rvk.docx"
    doc = Document()
    doc.add_paragraph("О СОСТОЯНИИ ЗДОРОВЬЯ ГРАЖДАНИНА №")
    doc.add_paragraph("Находился на обследовании")
    doc.add_paragraph("В 3 отделение КДП поступает")
    doc.add_paragraph("Жалобы")
    doc.add_paragraph("Анамнез жизни")
    doc.add_paragraph("Анамнез заболевания")
    doc.add_paragraph("Психический статус")
    doc.add_paragraph("Сомато-неврологический статус")
    doc.add_paragraph("Диагноз")
    doc.save(template)
    output = tmp_path / "rvk-out.docx"
    data = PatientData(
        fio="Иванов Иван Иванович", admission_date="01.09.2026", discharge_date="09.09.2026",
        admission_mode="первично", rvk_act_number="42", rvk_military_commissariat="Ленинский",
        diagnosis="F20.0 Параноидная шизофрения",
    )
    MedicalDocumentRenderer().render("rvk", template, output, data)
    text = "\n".join(p.text for p in Document(str(output)).paragraphs)
    assert "В 3 отделение КДП поступает первично" in text


def test_doctor_owned_discharge_postprocessing_enforces_mode_and_black_color(tmp_path: Path) -> None:
    template = tmp_path / "custom-discharge.docx"
    doc = Document()
    red = doc.styles.add_style("Doctor red template", WD_STYLE_TYPE.PARAGRAPH)
    red.font.color.rgb = RGBColor(255, 0, 0)
    period = doc.add_paragraph("Находилась на лечении в ГБУЗ НО «НКЦПЗ» диспансер №2 с 01.09.2026 по 09.09.2026")
    period.style = red
    doc.add_paragraph("В 3 отделение КДП поступает СТАРОЕ ЗНАЧЕНИЕ")
    doc.save(template)
    output = tmp_path / "custom-out.docx"
    case = PatientCase()
    case.update_from_pairs({
        "admission.date": "01.09.2026",
        "discharge.date": "09.09.2026",
        "admission.mode": "первично",
    })
    spec = DocumentTemplateSpec(
        id="discharge", button_label="Выписной эпикриз", template=str(template),
        role_id="discharge_epicrisis", required_fields=(), optional_fields=("admission.mode",),
    )
    render_template_to_docx(template_path=template, output_path=output, case=case, document=spec, strict=False)
    rendered = Document(str(output))
    text = "\n".join(p.text for p in rendered.paragraphs)
    assert "В 3 отделение КДП поступает первично" in text
    period = next(p for p in rendered.paragraphs if p.text.startswith("Находилась на лечении"))
    assert period.runs
    assert all(run.font.color.rgb == RGBColor(0, 0, 0) for run in period.runs)


def test_discharge_and_rvk_popup_contracts_require_primary_or_repeat_choice() -> None:
    discharge_source = Path("dialog_expert.py").read_text(encoding="utf-8")
    rvk_source = Path("dialog_document_details.py").read_text(encoding="utf-8")
    execution_source = Path("actions_creation_execution.py").read_text(encoding="utf-8")

    assert '("Поступает первично или повторно", "")' in discharge_source
    assert 'detail_fields.append("admission_mode")' in discharge_source
    assert 'normalize_admission_mode(value)' in discharge_source
    assert 'Выберите: первично или повторно.' in discharge_source

    assert 'text="Поступает первично или повторно?"' in rvk_source
    assert 'for idx, value in enumerate(("первично", "повторно"))' in rvk_source
    assert 'self.admission_mode_var.set(admission_mode)' in rvk_source
    assert 'Выберите, пациент поступает первично или повторно.' in rvk_source

    assert 'admission_mode = normalize_admission_mode(' in execution_source
    assert 'or not admission_mode' in execution_source


def test_primary_parser_recovers_explicit_admission_mode_without_guessing() -> None:
    from medical_parser import MedicalTextParser

    parser = MedicalTextParser()
    primary = parser.parse_text(
        "Первичный осмотр\n"
        "Ф.И.О.: Иванов Иван Иванович\n"
        "История болезни № 42\n"
        "Дата поступления: 01.09.2026\n"
        "В 3 отделение КДП поступает повторно\n"
        "Диагноз: F20.0 Параноидная шизофрения\n"
    )
    assert primary.admission_mode == "повторно"

    no_mode = parser.parse_text(
        "Первичный осмотр\n"
        "Ф.И.О.: Иванов Иван Иванович\n"
        "Дата поступления: 01.09.2026\n"
        "В 3 отделение КДП поступает добровольно\n"
        "Диагноз: F20.0 Параноидная шизофрения\n"
    )
    assert no_mode.admission_mode == ""


def test_service_requires_admission_mode_for_discharge_and_normalizes_synonym() -> None:
    import pytest
    from medical_service import MedicalDocumentService

    service = MedicalDocumentService()
    base = PatientData(
        fio="Иванов Иван Иванович",
        case_number="42",
        admission_date="01.09.2026",
        discharge_date="09.09.2026",
        diagnosis="F20.0 Параноидная шизофрения",
        treatment_plan="Терапия по схеме",
    )
    with pytest.raises(ValueError, match="первично или повторно"):
        service._validate_and_normalize_selected_data(base, ("discharge",))

    accepted = PatientData(**{**base.__dict__, "admission_mode": "первичный"})
    service._validate_and_normalize_selected_data(accepted, ("discharge",))
    assert accepted.admission_mode == "первично"
