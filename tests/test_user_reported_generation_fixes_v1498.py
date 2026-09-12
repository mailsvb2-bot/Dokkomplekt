from __future__ import annotations

from pathlib import Path
import re

from docx import Document

from diary_gender import adapt_text_to_patient_gender, detect_gender_from_patient_name
from diary_text_parser import extract_statuses_from_docx
from diary_text_selection import diary_filename_matches_diagnosis, find_diary_text_file_for_diagnosis
from medical_date_state import normalize_date_value
from medical_formatting import parse_date
from universal_fields import PatientCase
from universal_profiles import DocumentTemplateSpec
from universal_scanner import scan_docx
from universal_template_engine import render_template_to_docx


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def _write_lines(path: Path, *lines: str) -> Path:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def test_popup_date_contract_accepts_compact_six_digit_date() -> None:
    parsed = parse_date("090926")
    assert parsed is not None
    assert parsed.strftime("%d.%m.%Y") == "09.09.2026"
    assert normalize_date_value("090926") == "09.09.2026"


def test_short_diary_sentences_are_preserved_as_separate_sequence_entries(tmp_path: Path) -> None:
    source = _write_lines(tmp_path / "F32 тексты.docx", "Фон выравнивается.", "Активнее в режиме.")
    assert extract_statuses_from_docx(source) == ["Фон выравнивается.", "Активнее в режиме."]


def test_female_diary_text_adapts_reported_real_world_forms() -> None:
    gender = detect_gender_from_patient_name("Банина Екатерина Сергеевна")
    assert gender == "female"
    source = (
        "В отделение явился без сопровождения. На лечение пришёл самостоятельно. "
        "Остаётся вялым подавленным пассивным несобранным. "
        "Спокоен, вял, подавлен, несобран, малозаметен, бездеятелен, монотонен, выхолощен, холоден."
    )
    adapted, changed = adapt_text_to_patient_gender(source, gender)
    assert changed >= 10
    for male_form in (
        "явился", "пришёл", "вялым", "подавленным", "пассивным", "несобранным",
        "спокоен", "подавлен", "несобран", "малозаметен", "бездеятелен", "монотонен", "выхолощен", "холоден",
    ):
        assert re.search(rf"(?<![А-Яа-яЁё]){re.escape(male_form)}(?![А-Яа-яЁё])", adapted, flags=re.IGNORECASE) is None
    for female_form in ("явилась", "пришла", "вялой", "подавленной", "пассивной", "несобранной", "спокойна"):
        assert female_form.casefold() in adapted.casefold()


def test_diary_text_file_is_selected_by_diagnosis_filename_and_preserves_sequence(tmp_path: Path) -> None:
    texts = tmp_path / "тексты дневников"
    texts.mkdir()
    schizophrenia = _write_lines(
        texts / "шизофрения.docx",
        "Первый текст наблюдения достаточно длинный для дневника пациента.",
        "Одинаковый повторяющийся текст наблюдения достаточно длинный для дневника пациента.",
        "Одинаковый повторяющийся текст наблюдения достаточно длинный для дневника пациента.",
        "Последний текст наблюдения достаточно длинный для дневника пациента.",
    )
    _write_lines(texts / "депрессивный эпизод.docx", "Совершенно другой текст наблюдения для другого диагноза пациента.")

    selected = find_diary_text_file_for_diagnosis(texts, "F20.0 Параноидная шизофрения")
    assert selected == schizophrenia
    assert diary_filename_matches_diagnosis("F20.0 Параноидная шизофрения", "F20.0 шизофрения.docx")
    assert not diary_filename_matches_diagnosis("F20.0 Параноидная шизофрения", "F20.0 депрессивный эпизод.docx")
    statuses = extract_statuses_from_docx(selected)
    assert statuses == [
        "Первый текст наблюдения достаточно длинный для дневника пациента.",
        "Одинаковый повторяющийся текст наблюдения достаточно длинный для дневника пациента.",
        "Одинаковый повторяющийся текст наблюдения достаточно длинный для дневника пациента.",
        "Последний текст наблюдения достаточно длинный для дневника пациента.",
    ]


def test_scanner_keeps_labs_epi_and_epidemiology_semantically_separate(tmp_path: Path) -> None:
    source = _write_lines(
        tmp_path / "primary.docx",
        "Диагноз: F20.0 Параноидная шизофрения",
        "Эпидемиологический анамнез: Контактов с инфекционными больными не было.",
        "Результаты обследований:",
        "ОАК - Hb 140, лейкоциты 6.0",
        "ОАМ - белок не обнаружен",
        "ЭЭГ - без патологических изменений",
        "ЭПИ – Правильный текст ЭПИ",
        "За время лечения состояние улучшилось.",
        "Рекомендовано: наблюдение по месту жительства.",
    )
    best = scan_docx(source).best_matches()
    assert best["epidemiology"].value == "Контактов с инфекционными больными не было"
    assert best["epicrisis.text"].value == "Правильный текст ЭПИ"
    assert best["labs.results"].value == (
        "ОАК - Hb 140, лейкоциты 6.0\n"
        "ОАМ - белок не обнаружен\n"
        "ЭЭГ - без патологических изменений"
    )
    assert "labs.types" not in best
    assert best["treatment.result"].value == "состояние улучшилось"
    assert best["recommendations"].value == "наблюдение по месту жительства"


def test_discharge_generation_replaces_stale_labs_and_epi_with_canonical_values(tmp_path: Path) -> None:
    template = _write_lines(
        tmp_path / "discharge_template.docx",
        "ФИО: {{patient.fio}}",
        "Поступил: {{admission.date}}  Выписан: {{discharge.date}}",
        "Диагноз: {{diagnosis.main}}",
        "Результаты обследований:",
        "ОАК - ЧУЖОЙ АНАМНЕЗ ИЗ ШАБЛОНА",
        "ОАМ - ЧУЖОЙ АНАМНЕЗ ИЗ ШАБЛОНА",
        "RW - СТАРЫЕ ДАННЫЕ ДРУГОГО ПАЦИЕНТА",
        "ЭЭГ – ЧУЖОЙ АНАМНЕЗ ИЗ ШАБЛОНА",
        "ЭПИ – ЧУЖОЙ ЭПИ ИЗ ШАБЛОНА",
        "За время лечения: СТАРЫЙ РЕЗУЛЬТАТ",
    )
    output = tmp_path / "result.docx"
    case = PatientCase()
    case.update_from_pairs(
        {
            "patient.fio": "Банина Екатерина Сергеевна",
            "admission.date": "01.09.2026",
            "discharge.date": "09.09.2026",
            "diagnosis.main": "F20.0 Параноидная шизофрения",
            "labs.results": "ОАК - Hb 140\nОАМ - белок не обнаружен\nЭЭГ - без патологических изменений",
            "epicrisis.text": "Правильный текст ЭПИ",
            "treatment.result": "Состояние улучшилось.",
        },
        source_document="primary.docx",
    )
    spec = DocumentTemplateSpec(
        id="discharge",
        button_label="Выписной эпикриз",
        template=str(template),
        role_id="discharge_epicrisis",
        required_fields=("patient.fio", "admission.date", "discharge.date", "diagnosis.main"),
        optional_fields=("labs.results", "epicrisis.text", "treatment.result"),
    )
    render_template_to_docx(template_path=template, output_path=output, case=case, document=spec, strict=True)
    text = _docx_text(output)
    assert "ОАК - Hb 140\nОАМ - белок не обнаружен\nЭЭГ - без патологических изменений" in text
    assert "ЭПИ – Правильный текст ЭПИ" in text
    assert "За время лечения: Состояние улучшилось." in text
    for stale in ("ЧУЖОЙ АНАМНЕЗ", "СТАРЫЕ ДАННЫЕ ДРУГОГО ПАЦИЕНТА", "ЧУЖОЙ ЭПИ", "СТАРЫЙ РЕЗУЛЬТАТ"):
        assert stale not in text


def test_final_diary_text_is_fixed_independent_of_diagnosis_template(tmp_path: Path) -> None:
    from diary_batch import fill_diary_batch
    from diary_constants import FINAL_DIARY_TEXT

    final_lines: list[str] = []
    for index, (diagnosis_name, regular_text) in enumerate((
        ("шизофрения", "Промежуточный текст для шизофрении: контакту доступен, назначения принимает."),
        ("депрессивный эпизод", "Промежуточный текст для депрессии: фон настроения снижен, лечение получает."),
    )):
        source = _write_lines(tmp_path / f"{diagnosis_name}.docx", regular_text)
        result = fill_diary_batch(
            status_files=(source,), diary_files=(), output_dir=tmp_path / f"out-{index}",
            patient_name="Иванов Иван Иванович", admission_value="01.09.2026", discharge_value="03.09.2026",
            diary_day_offsets=(1, 2), force_final_diary=True, text_output=True,
        )
        text = _docx_text(Path(result.created_files[0]))
        expected = f"03.09.26 {FINAL_DIARY_TEXT}"
        assert expected in text
        final_lines.append(next(line for line in text.splitlines() if line.startswith("03.09.26 Состояние улучшилось")))

    assert final_lines[0] == final_lines[1]
