from __future__ import annotations

from pathlib import Path

from docx import Document

from medical_constants import DISCHARGE_RECOMMENDATION_TEXT
from medical_models import PatientData
from medical_renderer import MedicalDocumentRenderer
from universal_fields import PatientCase
from universal_profiles import DocumentTemplateSpec
from universal_template_engine import render_template_to_docx

BAD_RECOMMENDATION = (
    "Рекомендовано: По медицинским показаниям нуждается в оформлении "
    "академического отпуска на 2022-2023 учебный год"
)
STRAY_JOINT_COMPLAINT = "Пациентка предъявляет жалобы на апатию, отсутствие сил, нарушенный сон."
OTHER_STRAY_JOINT_COMPLAINT = "Пациент предъявляет жалобы на тревогу, слабость и ранние пробуждения."
FIXED_JOINT_BOILERPLATE = "Осмотр проведён совместно лечащим врачом и заведующим отделением."
FIXED_PATIENT_BOILERPLATE = "Пациент ознакомлен с рекомендациями, вопросы разъяснены."
FIXED_TREATMENT_BOILERPLATE = "Пациент проинформирован о плане лечения и возможных побочных эффектах."


def _text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def test_builtin_discharge_replaces_stale_template_recommendation(tmp_path: Path) -> None:
    template = tmp_path / "discharge.docx"
    doc = Document()
    doc.add_paragraph("Дата, время")
    doc.add_paragraph("Ф.И.О., г.р., зарегистрирован по адресу")
    doc.add_paragraph("Находилась на лечении в старом учреждении")
    doc.add_paragraph("В 3 отделение КДП поступает")
    doc.add_paragraph(BAD_RECOMMENDATION)
    doc.add_paragraph("Зав. отд. __________ Врач __________")
    doc.save(template)

    output = tmp_path / "discharge-out.docx"
    data = PatientData(
        fio="Иванова Анна Сергеевна",
        admission_date="01.09.2026",
        discharge_date="09.09.2026",
        admission_mode="первично",
    )
    MedicalDocumentRenderer().render("discharge", template, output, data)

    text = _text(output)
    assert DISCHARGE_RECOMMENDATION_TEXT in text
    assert "академического отпуска" not in text.lower()
    assert text.count(DISCHARGE_RECOMMENDATION_TEXT) == 1


def test_doctor_owned_discharge_replaces_stale_template_recommendation(tmp_path: Path) -> None:
    template = tmp_path / "custom-discharge.docx"
    doc = Document()
    doc.add_paragraph(BAD_RECOMMENDATION)
    doc.add_paragraph("Врач __________")
    doc.save(template)

    output = tmp_path / "custom-discharge-out.docx"
    case = PatientCase()
    spec = DocumentTemplateSpec(
        id="my-discharge",
        button_label="Моя выписка",
        template=str(template),
        role_id="discharge_epicrisis",
    )
    render_template_to_docx(
        template_path=template,
        output_path=output,
        case=case,
        document=spec,
        strict=False,
    )

    text = _text(output)
    assert DISCHARGE_RECOMMENDATION_TEXT in text
    assert "академического отпуска" not in text.lower()
    assert text.count(DISCHARGE_RECOMMENDATION_TEXT) == 1


def test_builtin_joint_exam_removes_stray_template_complaint(tmp_path: Path) -> None:
    template = tmp_path / "joint.docx"
    doc = Document()
    doc.add_paragraph("Совместный осмотр врачебной комиссией №")
    doc.add_paragraph("Жалобы при поступлении")
    doc.add_paragraph(STRAY_JOINT_COMPLAINT)
    doc.add_paragraph(OTHER_STRAY_JOINT_COMPLAINT)
    doc.add_paragraph(FIXED_JOINT_BOILERPLATE)
    doc.add_paragraph(FIXED_PATIENT_BOILERPLATE)
    doc.save(template)

    output = tmp_path / "joint-out.docx"
    data = PatientData(
        fio="Иванова Анна Сергеевна",
        commission_date="09.09.2026",
        commission_number="12",
        complaints="Жалоб не предъявляет.",
    )
    MedicalDocumentRenderer().render("commission", template, output, data)

    text = _text(output)
    assert STRAY_JOINT_COMPLAINT not in text
    assert OTHER_STRAY_JOINT_COMPLAINT not in text
    assert FIXED_JOINT_BOILERPLATE in text
    assert FIXED_PATIENT_BOILERPLATE in text
    assert "Жалобы при поступлении: Жалоб не предъявляет." in text


def test_doctor_owned_joint_exam_removes_stray_template_complaint(tmp_path: Path) -> None:
    template = tmp_path / "custom-joint.docx"
    doc = Document()
    doc.add_paragraph("Совместный осмотр")
    doc.add_paragraph(STRAY_JOINT_COMPLAINT)
    doc.add_paragraph(OTHER_STRAY_JOINT_COMPLAINT)
    doc.add_paragraph(FIXED_JOINT_BOILERPLATE)
    doc.add_paragraph(FIXED_PATIENT_BOILERPLATE)
    doc.save(template)

    output = tmp_path / "custom-joint-out.docx"
    case = PatientCase()
    spec = DocumentTemplateSpec(
        id="joint-custom",
        button_label="Переименованный документ",
        template=str(template),
        role_id="joint_medical_exam",
    )
    render_template_to_docx(
        template_path=template,
        output_path=output,
        case=case,
        document=spec,
        strict=False,
    )

    text = _text(output)
    assert STRAY_JOINT_COMPLAINT not in text
    assert OTHER_STRAY_JOINT_COMPLAINT not in text
    assert FIXED_JOINT_BOILERPLATE in text


def test_unrelated_same_diagnosis_cannot_legitimise_stale_complaint(tmp_path: Path) -> None:
    template = tmp_path / "same-code-stale-complaint.docx"
    doc = Document()
    stale = "Пациент предъявляет жалобы на выраженную тревогу; F20.0 указан в карте."
    doc.add_paragraph(stale)
    doc.save(template)

    output = tmp_path / "same-code-out.docx"
    case = PatientCase()
    case.update_from_pairs({
        "patient.fio": "Иванов Иван Иванович",
        "complaints": "Жалоб не предъявляет.",
        "diagnosis.main": "F20.0 Параноидная шизофрения",
    })
    spec = DocumentTemplateSpec(
        id="joint-code", button_label="Совместный осмотр", template=str(template),
        role_id="joint_medical_exam",
    )
    render_template_to_docx(
        template_path=template, output_path=output, case=case, document=spec, strict=False,
    )
    assert stale not in _text(output)


def test_gender_is_scoped_to_owned_phrase_not_static_template_boilerplate(tmp_path: Path) -> None:
    template = tmp_path / "female-discharge.docx"
    doc = Document()
    doc.add_paragraph("Находился на лечении в ГБУЗ НО «НКЦПЗ» диспансер №2 с 01.09.2026 по 09.09.2026")
    doc.add_paragraph(FIXED_PATIENT_BOILERPLATE)
    doc.save(template)

    output = tmp_path / "female-discharge-out.docx"
    case = PatientCase()
    case.update_from_pairs({
        "patient.fio": "Иванова Анна Сергеевна",
        "admission.date": "01.09.2026",
        "discharge.date": "09.09.2026",
    })
    spec = DocumentTemplateSpec(
        id="female-discharge", button_label="Выписной эпикриз", template=str(template),
        role_id="discharge_epicrisis",
    )
    render_template_to_docx(
        template_path=template, output_path=output, case=case, document=spec, strict=False,
    )
    text = _text(output)
    assert "Находилась на лечении" in text
    assert FIXED_PATIENT_BOILERPLATE in text
    assert "Пациентка ознакомлен" not in text


def test_doctor_owned_discharge_cleans_stale_footer_payload_and_keeps_footer_boilerplate(tmp_path: Path) -> None:
    template = tmp_path / "footer-discharge.docx"
    doc = Document()
    doc.add_paragraph("Выписной эпикриз")
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = BAD_RECOMMENDATION
    footer.add_paragraph("Документ сформирован лечащим подразделением.")
    doc.save(template)

    output = tmp_path / "footer-discharge-out.docx"
    case = PatientCase()
    spec = DocumentTemplateSpec(
        id="footer-discharge", button_label="Выписной эпикриз", template=str(template),
        role_id="discharge_epicrisis",
    )
    render_template_to_docx(
        template_path=template, output_path=output, case=case, document=spec, strict=False,
    )

    rendered = Document(str(output))
    footer_text = "\n".join(p.text for p in rendered.sections[0].footer.paragraphs)
    body_text = "\n".join(p.text for p in rendered.paragraphs)
    assert "академического отпуска" not in footer_text.lower()
    assert "Документ сформирован лечащим подразделением." in footer_text
    assert DISCHARGE_RECOMMENDATION_TEXT in body_text


def test_static_patient_boilerplate_with_document_date_is_not_deleted(tmp_path: Path) -> None:
    template = tmp_path / "dated-static-joint.docx"
    static = "Пациент ознакомлен с порядком обращения 09.09.2026, вопросы разъяснены."
    stale = "Пациент предъявляет жалобы на тревогу и ранние пробуждения."
    doc = Document()
    doc.add_paragraph("Совместный осмотр")
    doc.add_paragraph(static)
    doc.add_paragraph(stale)
    doc.save(template)

    output = tmp_path / "dated-static-joint-out.docx"
    case = PatientCase()
    case.update_from_pairs({"complaints": "Жалоб не предъявляет."})
    spec = DocumentTemplateSpec(
        id="dated-static-joint", button_label="Совместный осмотр", template=str(template),
        role_id="joint_medical_exam",
    )
    render_template_to_docx(
        template_path=template, output_path=output, case=case, document=spec, strict=False,
    )

    text = _text(output)
    assert static in text
    assert stale not in text



def test_patient_treatment_information_boilerplate_is_not_deleted(tmp_path: Path) -> None:
    template = tmp_path / "treatment-information-boilerplate.docx"
    stale = "Пациент предъявляет жалобы на тревогу и выраженную слабость."
    doc = Document()
    doc.add_paragraph("Совместный осмотр")
    doc.add_paragraph(FIXED_TREATMENT_BOILERPLATE)
    doc.add_paragraph(stale)
    doc.save(template)

    output = tmp_path / "treatment-information-boilerplate-out.docx"
    case = PatientCase()
    case.update_from_pairs({"complaints": "Жалоб не предъявляет."})
    spec = DocumentTemplateSpec(
        id="joint-treatment-info", button_label="Совместный осмотр", template=str(template),
        role_id="joint_medical_exam",
    )
    render_template_to_docx(
        template_path=template, output_path=output, case=case, document=spec, strict=False,
    )

    text = _text(output)
    assert FIXED_TREATMENT_BOILERPLATE in text
    assert stale not in text


def test_doctor_owned_discharge_preserves_canonical_template_recommendation(tmp_path: Path) -> None:
    template = tmp_path / "canonical-custom-discharge.docx"
    doc = Document()
    doc.add_paragraph(DISCHARGE_RECOMMENDATION_TEXT)
    doc.add_paragraph("Врач __________")
    doc.save(template)

    output = tmp_path / "canonical-custom-discharge-out.docx"
    case = PatientCase()
    spec = DocumentTemplateSpec(
        id="canonical-discharge", button_label="Выписной эпикриз", template=str(template),
        role_id="discharge_epicrisis",
    )
    render_template_to_docx(
        template_path=template, output_path=output, case=case, document=spec, strict=False,
    )

    text = _text(output)
    assert text.count(DISCHARGE_RECOMMENDATION_TEXT) == 1


def test_builtin_discharge_preserves_canonical_template_recommendation(tmp_path: Path) -> None:
    template = tmp_path / "canonical-builtin-discharge.docx"
    doc = Document()
    doc.add_paragraph("Дата, время")
    doc.add_paragraph("Ф.И.О., г.р., зарегистрирован по адресу")
    doc.add_paragraph("Находилась на лечении в старом учреждении")
    doc.add_paragraph("В 3 отделение КДП поступает")
    doc.add_paragraph(DISCHARGE_RECOMMENDATION_TEXT)
    doc.add_paragraph("Зав. отд. __________ Врач __________")
    doc.save(template)

    output = tmp_path / "canonical-builtin-discharge-out.docx"
    data = PatientData(
        fio="Иванова Анна Сергеевна",
        admission_date="01.09.2026",
        discharge_date="09.09.2026",
        admission_mode="первично",
    )
    MedicalDocumentRenderer().render("discharge", template, output, data)

    text = _text(output)
    assert text.count(DISCHARGE_RECOMMENDATION_TEXT) == 1
