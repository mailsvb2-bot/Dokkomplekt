"""Live GUI smoke for the production create/access/publication boundary.

Unlike the general doctor-emulation matrix this scenario deliberately enables
product access before pressing the real create command. It uses an isolated
license directory so it cannot consume a developer/user trial.

It also locks a proven July-2026 working behavior: a doctor-owned template must
still be created when all REQUIRED fields are present even if an explicitly
persisted OPTIONAL placeholder has no value. Later strict rendering treated
that optional blank as fatal and rolled the entire output transaction back.
"""

from __future__ import annotations

import os
import sys
import tkinter as tk
from dataclasses import replace
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT))

from doctor_sim import DoctorSim, make_docx  # noqa: E402


def _document_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def main() -> None:
    # False here deliberately chooses Word/DOCX in the real format popup.
    # This is the user-facing format required by the production workflow.
    sim = DoctorSim(answers={("askyesno", "Формат результата"): False})
    try:
        access_dir = sim.root_dir / "production-access"
        os.environ["DOKKOMPLEKT_TEST_DISABLE_PRODUCT_ACCESS"] = "0"
        os.environ["DOKKOMPLEKT_LICENSE_DIR"] = str(access_dir)

        primary = sim.root_dir / "production-access-primary.docx"
        make_docx(
            primary,
            [
                "История болезни № 777/26. Пациент: Иванов Иван Иванович, 1980 г.р.",
                "Дата поступления: 05.09.2026. Дата выписки: 08.09.2026.",
                "Диагноз: F32.1 Депрессивный эпизод средней степени.",
                "Лечение: сертралин 100 мг утром.",
            ],
        )
        texts = sim.root_dir / "production-access-texts-F32.docx"
        make_docx(texts, ["Состояние стабильное. Контактен, ориентирован."])

        sim.drop(primary)
        sim.answers["openfilenames"] = (str(texts),)
        sim.app.choose_status_files()
        sim.create_diaries()
        sim.pump(0.4)

        first_outputs = {path.resolve() for path in sim.outputs() if path.suffix.lower() == ".docx"}
        if sim.errors:
            raise AssertionError("real create command showed errors: " + " | ".join(sim.errors))
        if not first_outputs:
            raise AssertionError("real create command published no DOCX with product access enabled")

        # Historical doctor-template compatibility, through the current app's
        # real create command and the same product-access/output transaction.
        template = sim.root_dir / "doctor-optional-field.docx"
        template_doc = Document()
        template_doc.add_paragraph("Пациент: {{patient.fio}}")
        template_doc.add_paragraph("Диагноз: {{diagnosis.main}}")
        template_doc.add_paragraph("Больничный: {{expert.sick_leave_number}}")
        template_doc.save(template)

        from universal_template_engine import attach_template_to_pack
        from universal_profiles import save_document_pack
        from universal_main_documents import custom_kind

        pack = sim.app._load_or_create_universal_pack()
        spec, _target = attach_template_to_pack(
            pack,
            template,
            sim.app._universal_profile_path().parent,
            button_label="Историческая совместимость",
            document_id="historical_optional",
            role_id="discharge",
        )
        replacement = replace(
            spec,
            required_fields=("patient.fio", "diagnosis.main"),
            optional_fields=("expert.sick_leave_number",),
        )
        pack.documents = tuple(replacement if item.id == spec.id else item for item in pack.documents)
        save_document_pack(pack, sim.app._universal_profile_path())

        for variable in sim.app.output_vars.values():
            variable.set(False)
        # The production setup wizard normally creates the Tk variable when it
        # adds a doctor button. This smoke attaches a template after app startup,
        # so mirror only that UI-state registration here; generation itself stays
        # entirely production code.
        kind = custom_kind(spec.id)
        if kind not in sim.app.output_vars:
            variable = tk.BooleanVar(master=sim.tk_root, value=False)
            sim.app.output_vars[kind] = variable
            sim.app.custom_output_vars[kind] = variable
        sim.app.output_vars[kind].set(True)
        sim.app._on_output_toggle(kind)
        selected_after_toggle = sim.app.selected_custom_docs()
        if spec.id not in selected_after_toggle:
            raise AssertionError(
                "doctor-owned template selection was cancelled before create: "
                f"selected={selected_after_toggle}; popups={sim.popups[-8:]}"
            )
        creation_result = sim.app.create_selected_outputs(print_after=False)
        sim.pump(0.5)

        all_outputs = {path.resolve() for path in sim.outputs() if path.suffix.lower() == ".docx"}
        new_outputs = sorted(all_outputs - first_outputs)
        if sim.errors:
            raise AssertionError("doctor-template create command showed errors: " + " | ".join(sim.errors))
        if not creation_result:
            raise AssertionError(
                "doctor-owned top-level create command returned False: "
                f"selected={sim.app.selected_custom_docs()}; popups={sim.popups[-10:]}"
            )
        if not new_outputs:
            raise AssertionError(
                "doctor-owned template produced no published DOCX: "
                f"before={[str(p) for p in sorted(first_outputs)]}; "
                f"after={[str(p) for p in sorted(all_outputs)]}; popups={sim.popups[-10:]}"
            )
        doctor_output = next((path for path in new_outputs if "Историческая" in path.name), new_outputs[0])
        text = _document_text(doctor_output)
        if "Иванов Иван Иванович" not in text or "F32.1" not in text:
            raise AssertionError("primary-document values did not reach doctor-owned output: " + text[:400])
        if "{{expert.sick_leave_number}}" in text:
            raise AssertionError("empty optional placeholder survived in the generated DOCX")

        from product_access.native import NativeProductAccessManager

        state = NativeProductAccessManager(storage_dir=access_dir).current_state()
        if state.plan != "trial" or not state.active:
            raise AssertionError(f"unexpected product-access state: {state.plan}/{state.reason}")
        if state.documents_used_total_trial < len(all_outputs):
            raise AssertionError(
                f"published {len(all_outputs)} files but usage accounting is only {state.documents_used_total_trial}"
            )
        if not all(path.exists() and path.stat().st_size > 0 for path in all_outputs):
            raise AssertionError("one or more published outputs are missing/empty")

        print(
            "PRODUCTION CREATE ACCESS GUI SMOKE OK "
            f"files={len(all_outputs)} used={state.documents_used_total_trial} historical_template=1"
        )
    finally:
        sim.close()
        os.environ["DOKKOMPLEKT_TEST_DISABLE_PRODUCT_ACCESS"] = "1"
        os.environ.pop("DOKKOMPLEKT_LICENSE_DIR", None)


if __name__ == "__main__":
    main()
