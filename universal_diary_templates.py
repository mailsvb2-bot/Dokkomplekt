"""Profile support for doctor-owned diary templates.

Diary templates are special: many of them are table-based and do not contain
``{{field.id}}`` placeholders.  They are still valid doctor templates and must be
stored in the medpack with a confirmed date schedule.
"""

from __future__ import annotations

from dataclasses import replace
from diagnostic_logging import record_soft_exception
from pathlib import Path
from typing import Sequence

from diary_schedule import DiaryScheduleSpec, infer_diary_schedule_from_docx
from diary_text_parser import extract_statuses_from_docx
from universal_profiles import DocumentPack, DocumentTemplateSpec
from universal_template_engine import attach_template_to_pack, extract_template_placeholders
from docx import Document

UNIVERSAL_DIARY_TEMPLATE_LOCK_VERSION = "v1.1"
DIARY_TEMPLATES_CAN_BE_TABLE_BASED_WITHOUT_PLACEHOLDERS = True
DIARY_TEMPLATE_ERRORS_ARE_DIAGNOSTICALLY_LOGGED = True


def _normalized_cell_text(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("ё", "е").split())


def _table_looks_like_diary(table) -> bool:
    headers: list[str] = []
    rows = getattr(table, "rows", ()) or ()
    if not rows:
        return False
    for cell in rows[0].cells:
        headers.append(_normalized_cell_text(getattr(cell, "text", "")))
    joined = " ".join(headers)
    has_diary_column = any("дневник" in item or "наблюден" in item or "состояни" in item for item in headers) or "дневник" in joined
    has_hospital_day = any(("день" in item and "госпитал" in item) or "день госпитализации" in item for item in headers)
    has_date_columns = any("число" in item or "дата" in item for item in headers) or any("месяц" in item for item in headers)
    return bool(has_diary_column and (has_hospital_day or has_date_columns))


def looks_like_diary_template(path: str | Path) -> bool:
    candidate = Path(path).expanduser()
    if not candidate.exists() or candidate.suffix.lower() not in {".docx", ".docm"}:
        return False
    try:
        doc = Document(str(candidate))
        for table in doc.tables:
            if _table_looks_like_diary(table):
                return True
    except Exception as exc:
        record_soft_exception("universal_diary_templates.table_detection", exc)
    try:
        statuses = extract_statuses_from_docx(candidate)
        if statuses and "днев" in candidate.stem.lower():
            return True
    except Exception as exc:
        record_soft_exception("universal_diary_templates.status_detection", exc)
    try:
        placeholders = extract_template_placeholders(candidate)
        fields = {item.field_id for item in placeholders}
        return bool(fields & {"diary.entries", "diary.dates", "diary.schedule"})
    except Exception as exc:
        record_soft_exception("universal_diary_templates.placeholder_detection", exc)
        return False


def attach_diary_template_to_pack(
    pack: DocumentPack,
    template_path: str | Path,
    profile_dir: str | Path,
    *,
    button_label: str | None = None,
    document_id: str | None = None,
    schedule: DiaryScheduleSpec | None = None,
) -> tuple[DocumentTemplateSpec, Path, DiaryScheduleSpec]:
    inferred = schedule or infer_diary_schedule_from_docx([template_path])
    spec, copied_to = attach_template_to_pack(
        pack,
        template_path,
        profile_dir,
        button_label=button_label or "Дневники наблюдения",
        document_id=document_id or Path(template_path).stem,
        category="diaries",
        role_id="daily_diary",
        button_language="auto",
        source_language="auto",
        button_label_source="diary_template",
    )
    updated = replace(
        spec,
        category="diaries",
        role_id="daily_diary",
        description=(spec.description + " Дневниковый шаблон с подтверждаемым принципом дат.").strip(),
        diary_schedule=inferred.to_dict(),
    )
    pack.add_document(updated)
    return updated, copied_to, inferred


def diary_documents_with_hourly_mode(pack: DocumentPack) -> tuple[DocumentTemplateSpec, ...]:
    result: list[DocumentTemplateSpec] = []
    for document in pack.documents:
        if document.category != "diaries":
            continue
        spec = DiaryScheduleSpec.from_dict(getattr(document, "diary_schedule", None))
        if spec.has_hourly:
            result.append(document)
    return tuple(result)


def assert_universal_diary_template_lock() -> None:
    if UNIVERSAL_DIARY_TEMPLATE_LOCK_VERSION != "v1.1":
        raise AssertionError("Universal diary template lock changed unexpectedly")
    if not DIARY_TEMPLATES_CAN_BE_TABLE_BASED_WITHOUT_PLACEHOLDERS:
        raise AssertionError("Diary table templates must remain allowed without placeholders")
    if not DIARY_TEMPLATE_ERRORS_ARE_DIAGNOSTICALLY_LOGGED:
        raise AssertionError("Diary template detection errors must remain diagnosable")
